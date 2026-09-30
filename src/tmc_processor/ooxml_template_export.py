from __future__ import annotations
"""Direct OOXML writer for the Excel-authored four-leg v1 report template.

This module is isolated from production routing and never saves through openpyxl.
"""
import math, posixpath, shutil, zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable
from xml.dom import Node, minidom
import pandas as pd
from openpyxl.utils.cell import get_column_letter, range_boundaries
from .diagram import MOVEMENT_CODES
from .report_template import _mapped_label_value, _metadata_values, movement_values_by_code
from .template_write_plan import resolve_template_write_plan
MAIN_NS="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL_NS="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS="http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS="http://schemas.openxmlformats.org/package/2006/content-types"
CHART_NS="http://schemas.openxmlformats.org/drawingml/2006/chart"
VT_NS="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
XML_NS="http://www.w3.org/XML/1998/namespace"
WORKSHEET_REL=DOC_REL_NS+"/worksheet"
CALC_CHAIN_REL=DOC_REL_NS+"/calcChain"
WORKSHEET_CONTENT_TYPE="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"

@dataclass(frozen=True)
class OOXMLExportResult:
    output_path: Path
    changed_parts: tuple[str,...]
    added_parts: tuple[str,...]
    removed_parts: tuple[str,...]
    formula_cells_preserved: tuple[str,...]
    formula_cells_overwritten: tuple[str,...]

def _els(n,ns,tag): return n.getElementsByTagNameNS(ns,tag)
def _child(n,ns,tag):
    for x in n.childNodes:
        if x.nodeType==Node.ELEMENT_NODE and x.namespaceURI==ns and x.localName==tag: return x
    return None
def _text(n):
    if n is None: return ""
    return "".join(x.data for x in n.childNodes if x.nodeType in (Node.TEXT_NODE,Node.CDATA_SECTION_NODE))
def _remove(n,ns,*tags):
    for x in list(n.childNodes):
        if x.nodeType==Node.ELEMENT_NODE and x.namespaceURI==ns and x.localName in tags:
            n.removeChild(x); x.unlink()
def _xml(doc): return doc.toxml(encoding="utf-8")
def _target(base,target): return target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(base),target))
def _sheet_parts(z):
    wb=minidom.parseString(z.read("xl/workbook.xml")); rel=minidom.parseString(z.read("xl/_rels/workbook.xml.rels"))
    targets={x.getAttribute("Id"):x.getAttribute("Target") for x in _els(rel,PKG_REL_NS,"Relationship") if x.getAttribute("TargetMode").casefold()!="external"}
    out={}
    for s in _els(wb,MAIN_NS,"sheet"):
        rid=s.getAttributeNS(DOC_REL_NS,"id"); part=_target("xl/workbook.xml",targets.get(rid,""))
        if not part or part not in z.namelist(): raise ValueError(f"Worksheet relationship/part missing: {s.getAttribute('name')}")
        out[s.getAttribute("name")]=part
    return out
def _shared(z):
    if "xl/sharedStrings.xml" not in z.namelist(): return []
    root=minidom.parseString(z.read("xl/sharedStrings.xml"))
    return ["".join(_text(t) for t in _els(si,MAIN_NS,"t")) for si in _els(root,MAIN_NS,"si")]
def _cellmap(doc): return {c.getAttribute("r"):c for c in _els(doc,MAIN_NS,"c")}
def _cellvalue(c,strings):
    if c is None: return None
    typ=c.getAttribute("t")
    if typ=="inlineStr": return "".join(_text(t) for t in _els(_child(c,MAIN_NS,"is"),MAIN_NS,"t"))
    v=_text(_child(c,MAIN_NS,"v"))
    if typ=="s": return strings[int(v)] if v else ""
    if typ=="b": return v=="1"
    if typ=="str": return v
    if not v: return None
    try:
        num=float(v); return int(num) if num.is_integer() else num
    except ValueError: return v
def _scalar(v):
    if v is None or v is pd.NA or v is pd.NaT: return ""
    if isinstance(v,pd.Timestamp): return v.to_pydatetime()
    try:
        if not isinstance(v,(str,bytes,list,tuple,dict,date,datetime,time)) and pd.isna(v): return ""
    except (TypeError,ValueError): pass
    if hasattr(v,"item") and callable(v.item):
        try: v=v.item()
        except (ValueError,TypeError): pass
    return v
def _date_serial(v): return (v.replace(tzinfo=None)-datetime(1899,12,30)).total_seconds()/86400
def _set_value(doc,cell,value,formula=False):
    value=_scalar(value); _remove(cell,MAIN_NS,"f","v","is")
    if cell.hasAttribute("t"): cell.removeAttribute("t")
    if formula:
        f=doc.createElementNS(MAIN_NS,"f"); text=str(value); f.appendChild(doc.createTextNode(text[1:] if text.startswith("=") else text)); cell.appendChild(f); cell.appendChild(doc.createElementNS(MAIN_NS,"v")); return
    if isinstance(value,str):
        cell.setAttribute("t","inlineStr"); inline=doc.createElementNS(MAIN_NS,"is"); t=doc.createElementNS(MAIN_NS,"t")
        if value[:1].isspace() or value[-1:].isspace(): t.setAttributeNS(XML_NS,"xml:space","preserve")
        if value: t.appendChild(doc.createTextNode(value))
        inline.appendChild(t); cell.appendChild(inline)
    elif isinstance(value,bool):
        cell.setAttribute("t","b"); v=doc.createElementNS(MAIN_NS,"v"); v.appendChild(doc.createTextNode("1" if value else "0")); cell.appendChild(v)
    elif isinstance(value,datetime):
        v=doc.createElementNS(MAIN_NS,"v"); v.appendChild(doc.createTextNode(format(_date_serial(value),".15g"))); cell.appendChild(v)
    elif isinstance(value,date): _set_value(doc,cell,value.isoformat())
    elif isinstance(value,time): _set_value(doc,cell,value.strftime("%H:%M"))
    elif isinstance(value,(int,float)):
        if isinstance(value,float) and not math.isfinite(value): _set_value(doc,cell,""); return
        v=doc.createElementNS(MAIN_NS,"v"); v.appendChild(doc.createTextNode(str(value))); cell.appendChild(v)
    else: _set_value(doc,cell,str(value))
def _row(doc,row_num):
    data=_els(doc,MAIN_NS,"sheetData")[0]
    for r in _els(data,MAIN_NS,"row"):
        if int(r.getAttribute("r"))==row_num: return r
    r=doc.createElementNS(MAIN_NS,"row"); r.setAttribute("r",str(row_num)); inserted=False
    for old in _els(data,MAIN_NS,"row"):
        if int(old.getAttribute("r"))>row_num: data.insertBefore(r,old); inserted=True; break
    if not inserted: data.appendChild(r)
    return r
def _ensure_cell(doc,cells,ref):
    if ref in cells: return cells[ref]
    col,row,_,_=range_boundaries(ref); r=_row(doc,row); cell=doc.createElementNS(MAIN_NS,"c"); cell.setAttribute("r",ref); inserted=False
    for old in _els(r,MAIN_NS,"c"):
        if range_boundaries(old.getAttribute("r"))[0]>col: r.insertBefore(cell,old); inserted=True; break
    if not inserted: r.appendChild(cell)
    cells[ref]=cell; return cell
def _in_range(ref,ranges):
    col,row,_,_=range_boundaries(ref)
    for rg in ranges:
        try: c1,r1,c2,r2=range_boundaries(str(rg).split("!")[-1].replace("$",""))
        except ValueError: continue
        if c1<=col<=c2 and r1<=row<=r2: return True
    return False
def _put_summary(doc,cells,ref,value,mapping,*,intent,force=False,preserve=False):
    if not ref: return False
    writable=list(mapping.get("writable_ranges") or []); overwrite=list(mapping.get("formula_overwrite_ranges") or []); protected=list(mapping.get("protected_formula_ranges") or [])
    allowed=_in_range(ref,writable) or _in_range(ref,overwrite) or force
    if not allowed:
        if preserve and _in_range(ref,protected): return False
        raise ValueError(f"Unsafe OOXML write rejected at Summary!{ref} ({intent}).")
    cell=_ensure_cell(doc,cells,ref)
    if _child(cell,MAIN_NS,"f") is not None and not _in_range(ref,overwrite) and not force:
        if preserve and _in_range(ref,protected): return False
        raise ValueError(f"Unsafe OOXML formula overwrite rejected at Summary!{ref} ({intent}).")
    _set_value(doc,cell,value,formula=isinstance(value,str) and value.startswith("=")); return True

def _normal_time(v): return str(v or "").strip().replace(".",":").lower()
def _colnum(s):
    n=0
    for ch in s.upper():
        if "A"<=ch<="Z": n=n*26+ord(ch)-64
    return n
def _column_for(frame,key):
    if not len(frame.columns): return None
    if key=="time": return str(frame.columns[0])
    wanted={"total_pcu":"Total (PCU)","total_vehicles":"Total (à¸„à¸±à¸™)"}.get(key,key)
    if key=="total_vehicles": wanted="Total (คัน)"
    if wanted in frame.columns: return str(wanted)
    return {str(c).strip().lower():str(c) for c in frame.columns}.get(str(wanted).strip().lower())
def _source_rows(frame):
    col=_column_for(frame,"time"); labels={}; total=None
    if not col: return labels,total
    for _,row in frame.iterrows():
        key=_normal_time(row[col])
        if key in {"รวม","total"}: total=row
        elif key: labels[key]=row
    return labels,total
def _write_tables(doc,cells,z,mapping,data):
    strings=_shared(z); summary=minidom.parseString(z.read(_sheet_parts(z)[mapping.get("template_sheet","Summary")]))
    template_cells=_cellmap(summary)
    sheets=data.get("sheets") or {}
    hourly_vehicle_class=data.get("hourly_vehicle_class")
    if hourly_vehicle_class is None:
        hourly_vehicle_class=sheets.get("Hourly_Vehicle_Class",pd.DataFrame())
    specs=((mapping.get("hourly_movement_table",{}),data.get("hourly_movement_pcu",pd.DataFrame())),(mapping.get("hourly_vehicle_class_table",{}),hourly_vehicle_class))
    for table,frame in specs:
        if not table: continue
        head=int(table["header_row"]); first=int(table["first_data_row"]); total=int(table.get("total_row") or first+max(len(frame),1)-1); columns=table.get("columns",{}); time_col=table.get("time_column")
        labels={row:_cellvalue(template_cells.get(f"{time_col}{row}"),strings) for row in range(first,total+1)} if time_col else {}
        by_label,source_total=_source_rows(frame)
        for key,col in columns.items():
            source_col=_column_for(frame,key)
            _put_summary(doc,cells,f"{col}{head}",source_col or key,mapping,intent=f"table.header.{key}",preserve=True)
            for row in range(first,total+1):
                if labels:
                    source=source_total if row==total else by_label.get(_normal_time(labels[row]))
                    if key=="time": value=labels[row]
                    elif source is not None and source_col and source_col in frame.columns: value=source[source_col]
                    else: value=0
                else:
                    offset=row-first; value=frame.iloc[offset][source_col] if offset<len(frame) and source_col and source_col in frame.columns else ""
                _put_summary(doc,cells,f"{col}{row}",value,mapping,intent=f"table.value.{key}",preserve=True)
def _write_chart_sources(doc,cells,mapping,chart_data):
    anchors=mapping.get("chart_anchors",{})
    for src_key,data_key in (("native_hourly_chart_source","hourly_pcu"),("native_vehicle_composition_chart_source","vehicle_composition")):
        src=anchors.get(src_key,{}) or {}; data=chart_data.get(data_key,{})
        for range_key in ("categories","values"):
            area=src.get(range_key)
            if not area: continue
            c1,r1,c2,r2=range_boundaries(area.replace("$","")); vals=list(data.get(range_key,[])); cap=(c2-c1+1)*(r2-r1+1)
            if len(vals)>cap: raise ValueError(f"Chart data exceeds mapped range {area}.")
            for offset in range(cap):
                row=r1+offset//(c2-c1+1); col=c1+offset%(c2-c1+1); ref=f"{get_column_letter(col)}{row}"
                _put_summary(doc,cells,ref,vals[offset] if offset<len(vals) else "",mapping,intent=f"chart.{data_key}.{range_key}",preserve=True)
def _clear_formula_caches(doc):
    for c in _els(doc,MAIN_NS,"c"):
        if _child(c,MAIN_NS,"f") is not None:
            if c.hasAttribute("t"):
                c.removeAttribute("t")
            v=_child(c,MAIN_NS,"v")
            if v is not None:
                for child in list(v.childNodes): v.removeChild(child); child.unlink()
def _set_formula_cache(doc,cells,ref,value):
    cell=cells.get(ref)
    if cell is None or _child(cell,MAIN_NS,"f") is None: return
    if isinstance(value,str): cell.setAttribute("t","str")
    elif cell.hasAttribute("t"): cell.removeAttribute("t")
    cache=_child(cell,MAIN_NS,"v")
    if cache is None:
        cache=doc.createElementNS(MAIN_NS,"v"); cell.appendChild(cache)
    for child in list(cache.childNodes): cache.removeChild(child); child.unlink()
    cache.appendChild(doc.createTextNode(str(value)))


def _cache_required_template_sources(doc,cells,strings,template_map,plan):
    """Cache only the authoritative template's fixed source/rollup formulas."""
    hourly=template_map["hourly_movement_table"]
    first=int(hourly["first_data_row"]); last=int(hourly["last_data_row"]); total=int(hourly["total_row"])
    diagram={item.code:item for item in plan.diagram_rows}
    support=dict(plan.support_sheets).get("Hourly_Movement_PCU")
    if support is None: raise ValueError("Hourly_Movement_PCU support sheet is required for native total caches.")
    movement_totals=[]
    for key,column in hourly["columns"].items():
        if key=="time": continue
        ref=f"{column}{total}"
        cell=cells.get(ref); formula=_child(cell,MAIN_NS,"f") if cell is not None else None
        if formula is None or (_text(formula) and _text(formula)!=f"SUM({column}{first}:{column}{last})") or (not _text(formula) and formula.getAttribute("t")!="shared"):
            raise ValueError(f"Native hourly total formula changed at Summary!{ref}.")
        values=[_cellvalue(cells.get(f"{column}{row}"),strings) for row in range(first,last+1)]
        if not all(isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(float(value)) for value in values):
            raise ValueError(f"Non-numeric hourly source for Summary!{ref}.")
        amount=sum(values)
        if key!="Total":
            if key not in diagram or key not in support.columns or not math.isclose(amount,diagram[key].total_cache,rel_tol=1e-9,abs_tol=1e-9):
                raise ValueError(f"Hourly movement total disagrees with Diagram_Data for {key}.")
            movement_totals.append(amount)
        elif not math.isclose(amount,sum(movement_totals),rel_tol=1e-9,abs_tol=1e-9):
            raise ValueError("Grand hourly total does not equal the 16 movement totals.")
        if key in support.columns and not math.isclose(amount,float(support.iloc[-1][key]),rel_tol=1e-9,abs_tol=1e-9):
            raise ValueError(f"Hourly movement total disagrees with support-sheet total for {key}.")
        _set_formula_cache(doc,cells,ref,amount)

    rollups={
        "F30":("G19:G22","J16:M16","O21:O24","J27:M27"),
        "F31":("J15:M15","F19:F22","J28:M28","P21:P24"),
        "F32":("J14:M14","E19:E22","J29:M29","Q21:Q24"),
    }
    for ref,areas in rollups.items():
        cell=cells.get(ref); formula=_child(cell,MAIN_NS,"f") if cell is not None else None
        if formula is None or _text(formula)!=f"SUM({','.join(areas)})":
            raise ValueError(f"Native report rollup formula changed at Summary!{ref}.")
        values=[_cellvalue(cells.get(source),strings) for area in areas for source in _range_refs(area)]
        if not all(isinstance(value,(int,float)) and not isinstance(value,bool) for value in values):
            raise ValueError(f"Missing movement source cache for Summary!{ref}.")
        amount=sum(values)
        metric={"F30":"am_cache","F31":"pm_cache","F32":"total_cache"}[ref]
        if not math.isclose(amount,sum(getattr(item,metric) for item in plan.diagram_rows),rel_tol=1e-9,abs_tol=1e-9):
            raise ValueError(f"Report rollup does not reconcile with all 16 movements at Summary!{ref}.")
        _set_formula_cache(doc,cells,ref,amount)

    title=cells.get("C8"); title_formula=_child(title,MAIN_NS,"f") if title is not None else None
    if title_formula is None or _text(title_formula)!="E5": raise ValueError("Native diagram title formula changed at Summary!C8.")
    _set_formula_cache(doc,cells,"C8",_cellvalue(cells.get("E5"),strings) or "")

    # =+K5 can be cached safely only for a numeric or blank K5. The product's
    # survey_date_text may be arbitrary text, whose Excel coercion is locale
    # dependent; do not manufacture a date serial or an error cache here.
    date=cells.get("C9"); date_formula=_child(date,MAIN_NS,"f") if date is not None else None
    if date_formula is None or _text(date_formula)!="+K5": raise ValueError("Native diagram date formula changed at Summary!C9.")
    source=_cellvalue(cells.get("K5"),strings)
    if source is None: _set_formula_cache(doc,cells,"C9",0)
    elif isinstance(source,(int,float)) and not isinstance(source,bool): _set_formula_cache(doc,cells,"C9",source)
def _range_refs(area):
    c1,r1,c2,r2=range_boundaries(area.replace("$","")); refs=[]
    for row in range(r1,r2+1):
        for col in range(c1,c2+1): refs.append(f"{get_column_letter(col)}{row}")
    return refs
def _cache_visible_report_formulas(doc,cells,strings,template_map,chart_source_data):
    hourly=template_map["hourly_movement_table"]; vehicle=template_map["hourly_vehicle_class_table"]
    first=int(vehicle["first_data_row"]); last=int(vehicle["last_data_row"]); total=int(vehicle["total_row"])
    source_first=int(hourly["first_data_row"]); source_col=hourly["columns"]["Total"]
    target_pcu_col=vehicle["columns"]["total_pcu"]
    if last-first!=int(hourly["last_data_row"])-source_first:
        raise ValueError("Hourly vehicle and movement Summary tables have incompatible row counts.")
    for offset,row in enumerate(range(first,last+1)):
        value=_cellvalue(cells.get(f"{source_col}{source_first+offset}"),strings)
        if not isinstance(value,(int,float)):
            raise ValueError(f"Expected numeric chart PCU source at Summary!{source_col}{source_first+offset}.")
        _set_formula_cache(doc,cells,f"{target_pcu_col}{row}",value)
    total_columns=[col for key,col in vehicle["columns"].items() if key!="time"]
    totals={}
    for col in total_columns:
        values=[_cellvalue(cells.get(f"{col}{row}"),strings) for row in range(first,last+1)]
        if not all(isinstance(value,(int,float)) for value in values):
            raise ValueError(f"Expected numeric vehicle Summary inputs in {col}{first}:{col}{last}.")
        totals[col]=sum(values)
        _set_formula_cache(doc,cells,f"{col}{total}",totals[col])
    class_columns=[col for key,col in vehicle["columns"].items() if key not in {"time","total_pcu","total_vehicles"}]
    class_headers={_cellvalue(cells.get(f"{col}{int(vehicle['header_row'])}"),strings):col for col in class_columns}
    anchors=template_map.get("chart_anchors",{}).get("native_vehicle_composition_chart_source",{})
    category_refs=_range_refs(anchors.get("categories","")); value_refs=_range_refs(anchors.get("values",""))
    categories=[_cellvalue(cells.get(ref),strings) for ref in category_refs]
    if not category_refs or len(category_refs)!=len(value_refs) or len(categories)!=len(class_columns):
        raise ValueError("Native vehicle chart ranges do not match the vehicle-class Summary columns.")
    denominator=sum(totals[col] for col in class_columns)
    if denominator<=0: raise ValueError("Vehicle-composition Summary totals must have a positive denominator.")
    expected=[]
    for category,value_ref in zip(categories,value_refs):
        col=class_headers.get(category)
        if col is None: raise ValueError(f"Vehicle chart category {category!r} has no Summary total column.")
        class_total=totals[col]; expected.append(class_total/denominator)
        column_letters="".join(char for char in category_refs[len(expected)-1] if char.isalpha())
        category_row=int("".join(char for char in category_refs[len(expected)-1] if char.isdigit()))
        _set_formula_cache(doc,cells,f"{column_letters}{category_row-1}",class_total)
        _set_formula_cache(doc,cells,value_ref,expected[-1])
    supplied=list((chart_source_data.get("vehicle_composition") or {}).get("values",[]))
    supplied_categories=list((chart_source_data.get("vehicle_composition") or {}).get("categories",[]))
    if supplied_categories!=categories or len(supplied)!=len(expected) or any(not math.isclose(float(a),b,rel_tol=1e-9,abs_tol=1e-12) for a,b in zip(supplied,expected)):
        raise ValueError("Native vehicle chart source values do not reconcile with the Summary vehicle-class totals.")
def _cache_summary_movement_values(doc,cells,strings,template_map,plan,movement_values):
    movements=template_map.get("movement_diagram_cells",{}).get("diagram_movements",{})
    columns=template_map["hourly_movement_table"]["columns"]
    peak_rows={peak.period:peak.worksheet_row for peak in plan.peak_formula_rows}
    for code,info in movements.items():
        metrics=(
            ("total_12_hour_cell","total_12_hour"),
            ("pm_peak_hour_cell","pm_peak"),
            ("am_peak_hour_cell","am_peak"),
        )
        for cell_key,value_key in metrics:
            ref=info.get(cell_key)
            if not ref: continue
            if code not in movement_values: raise ValueError(f"Movement diagram value is missing for {code}.")
            if _child(cells.get(ref),MAIN_NS,"f") is None:
                raise ValueError(f"Mapped movement display cell Summary!{ref} does not retain its template formula.")
            expected=movement_values[code][value_key]
            if value_key in {"am_peak","pm_peak"}:
                period=value_key[:2]
                source=f"{columns[code]}{peak_rows[period]}"
                selected=_cellvalue(cells.get(source),strings)
                if not isinstance(selected,(int,float)) or not math.isclose(selected,expected,rel_tol=1e-9,abs_tol=1e-9):
                    raise ValueError(f"Selected Summary Peak row disagrees with Diagram_Data for {code} {period.upper()}.")
                expected=selected
            _set_formula_cache(doc,cells,ref,expected)

def _validate_movement_payload(data,metadata,template_map):
    sheets=data.get("sheets") or {}
    required=("Mapping","Normalized_Data","Movement_Summary")
    present=[name for name in required if name in sheets]
    if not present: return
    if len(present)!=len(required): raise ValueError(f"Movement completeness validation requires support sheets {required}.")
    mapping=sheets["Mapping"]; normalized=sheets["Normalized_Data"]; movement_summary=sheets["Movement_Summary"]
    hourly=data.get("hourly_movement_pcu")
    if hourly is None or hourly.empty: raise ValueError("Movement completeness validation requires hourly movement payload data.")
    map_columns=set(map(str,mapping.columns)); norm_columns=set(map(str,normalized.columns)); summary_columns=set(map(str,movement_summary.columns))
    map_required={"from_leg","to_leg","turn_type","movement_code","include_in_report"}
    norm_required={"from_leg","to_leg","turn_type","output_movement_code","include_in_report","include_in_peak","time_start","pcu"}
    summary_required={"from_leg","to_leg","turn_type","movement_code","include_in_report","pcu"}
    if not map_required<=map_columns or not norm_required<=norm_columns or not summary_required<=summary_columns:
        raise ValueError("Movement completeness support sheets are missing semantic identity/value columns.")
    def included(value): return bool(value) and str(value).strip().casefold() not in {"false","0","no"}
    def identity(row,code_key): return (str(row["from_leg"]),str(row["to_leg"]),str(row["turn_type"]),str(row[code_key]))
    expected={identity(row,"movement_code") for _,row in mapping.iterrows() if included(row["include_in_report"])}
    actual_normalized={identity(row,"output_movement_code") for _,row in normalized.iterrows() if included(row["include_in_report"])}
    actual_summary={identity(row,"movement_code") for _,row in movement_summary.iterrows() if included(row["include_in_report"])}
    if not expected or expected!=actual_normalized or expected!=actual_summary:
        raise ValueError(f"Mapped, normalized, and movement-summary identity sets differ: expected={len(expected)}, normalized={len(actual_normalized)}, summary={len(actual_summary)}.")
    code_identities={}
    for from_leg,to_leg,turn_type,code in expected:
        if code in code_identities and code_identities[code]!=(from_leg,to_leg,turn_type):
            raise ValueError(f"Movement code {code!r} identifies multiple semantic movements.")
        code_identities[code]=(from_leg,to_leg,turn_type)
    diagram=template_map.get("movement_diagram_cells",{}).get("diagram_movements",{})
    if not set(code_identities)<=set(diagram): raise ValueError(f"Expected movement codes have no mapped diagram slots: {sorted(set(code_identities)-set(diagram))}.")
    refs=[]
    for code in code_identities:
        info=diagram[code]
        refs.extend(info.get(key) for key in ("header_cell","total_12_hour_cell","pm_peak_hour_cell","am_peak_hour_cell") if info.get(key))
    if len(refs)!=len(set(refs)): raise ValueError("Multiple movement identities map to the same visible diagram cell.")

    values=movement_values_by_code(hourly,list(code_identities),metadata)
    time_column=str(hourly.columns[0])
    start_rows={str(row[time_column]).strip().split("-")[0].strip():row for _,row in hourly.iterrows() if str(row[time_column]).strip().casefold()!="total"}
    def in_effective_hour(value,start):
        try:
            label=str(value).strip()[:5]
            hour,minute=(int(part) for part in label.split(":"))
            start_hour,start_minute=(int(part) for part in start.split(":"))
            actual=hour*60+minute; lower=start_hour*60+start_minute
        except (TypeError,ValueError):
            return False
        return lower<=actual<lower+60
    for period in ("am","pm"):
        start=str(metadata.get(f"{period}_peak_start") or "").strip()
        if start not in start_rows: raise ValueError(f"Effective {period.upper()} Peak {start!r} is absent from hourly movement payload.")
        payload_row=start_rows[start]
        expected_total=0.0
        for from_leg,to_leg,turn_type,code in expected:
            norm_values=normalized[(normalized["from_leg"].astype(str)==from_leg)&(normalized["to_leg"].astype(str)==to_leg)&(normalized["turn_type"].astype(str)==turn_type)&(normalized["output_movement_code"].astype(str)==code)&normalized["include_in_report"].map(included)&normalized["include_in_peak"].map(included)&normalized["time_start"].map(lambda value:in_effective_hour(value,start))]
            expected_value=sum(float(value or 0) for value in norm_values["pcu"])
            payload_value=float(payload_row[code] or 0)
            if not math.isclose(expected_value,payload_value,rel_tol=1e-9,abs_tol=1e-9):
                raise ValueError(f"{period.upper()} payload mismatch for {from_leg}->{to_leg} {turn_type} ({code}): normalized={expected_value}, hourly={payload_value}.")
            values[code][f"{period}_peak"]=payload_value
            expected_total+=expected_value
        total_column=next((column for column in hourly.columns if str(column).strip().casefold()=="total"),None)
        if total_column is None or not math.isclose(expected_total,float(payload_row[total_column] or 0),rel_tol=1e-9,abs_tol=1e-9):
            raise ValueError(f"{period.upper()} mapped movement sum does not reconcile to hourly total.")
    summary_by_identity={identity(row,"movement_code"):float(row["pcu"] or 0) for _,row in movement_summary.iterrows() if included(row["include_in_report"])}
    for movement in expected:
        code=movement[3]
        total=values[code]["total_12_hour"]
        if not math.isclose(total,summary_by_identity[movement],rel_tol=1e-9,abs_tol=1e-9):
            raise ValueError(f"12-hour movement total mismatch for {movement}.")
def _diagram_formula(kind,ref,rows,columns):
    last=get_column_letter(max(columns,2)); data=f"'Hourly_Movement_PCU'!$B$2:${last}${rows}"; times=f"'Hourly_Movement_PCU'!$A$2:$A${rows}"; headers=f"'Hourly_Movement_PCU'!$B$1:${last}$1"; match=f"MATCH({ref},{headers},0)"
    if kind=="total": return f"=IFERROR(INDEX({data},ROWS({times}),{match}),0)"
    peak="'Peak_PHF'!$H$2" if kind=="pm" else "'Peak_PHF'!$D$2"; key=f'IF(ISNUMBER({peak}),TEXT({peak},"hh:mm"),LEFT({peak},5))&"*"'
    return f"=IFERROR(INDEX({data},MATCH({key},{times},0),{match}),0)"
def _worksheet_shell(doc,cols,rows):
    root=doc.documentElement; root.setAttribute("xmlns",MAIN_NS)
    dimension=doc.createElementNS(MAIN_NS,"dimension"); dimension.setAttribute("ref",f"A1:{get_column_letter(cols)}{rows}"); root.appendChild(dimension)
    views=doc.createElementNS(MAIN_NS,"sheetViews"); view=doc.createElementNS(MAIN_NS,"sheetView"); view.setAttribute("workbookViewId","0"); views.appendChild(view); root.appendChild(views)
    fmt=doc.createElementNS(MAIN_NS,"sheetFormatPr"); fmt.setAttribute("defaultRowHeight","15"); root.appendChild(fmt)
def _dataframe_sheet(frame,header_style):
    doc=minidom.getDOMImplementation().createDocument(MAIN_NS,"worksheet",None); data=doc.createElementNS(MAIN_NS,"sheetData")
    rows=[[str(c) for c in frame.columns]]+[[v for v in row] for row in frame.itertuples(index=False,name=None)]
    if len(frame.columns)==0: rows=[[""]]
    ncol=max(max((len(r) for r in rows),default=1),1); _worksheet_shell(doc,ncol,max(len(rows),1))
    cols=doc.createElementNS(MAIN_NS,"cols")
    for i in range(ncol):
        vals=[r[i] for r in rows if i<len(r)]; width=min(255,max(10,max((len(str(_scalar(v))) for v in vals),default=8)+2)); col=doc.createElementNS(MAIN_NS,"col")
        col.setAttribute("min",str(i+1)); col.setAttribute("max",str(i+1)); col.setAttribute("width",str(width)); col.setAttribute("customWidth","1"); cols.appendChild(col)
    doc.documentElement.insertBefore(cols,_child(doc.documentElement,MAIN_NS,"sheetData"))
    for ri,values in enumerate(rows,1):
        row=doc.createElementNS(MAIN_NS,"row"); row.setAttribute("r",str(ri))
        for ci,value in enumerate(values,1):
            cell=doc.createElementNS(MAIN_NS,"c"); cell.setAttribute("r",f"{get_column_letter(ci)}{ri}")
            if ri==1: cell.setAttribute("s",str(header_style))
            _set_value(doc,cell,value); row.appendChild(cell)
        data.appendChild(row)
    doc.documentElement.appendChild(data); return _xml(doc)
def _diagram_sheet(frame,codes,header_style,metadata):
    doc=minidom.getDOMImplementation().createDocument(MAIN_NS,"worksheet",None); _worksheet_shell(doc,4,len(codes)+1); data=doc.createElementNS(MAIN_NS,"sheetData"); headers=("movement_code","total_pcu","pm_peak_pcu","am_peak_pcu"); rows=max(len(frame)+1,2); cols=max(len(frame.columns),2)
    content=[headers]+[(code,_diagram_formula("total",f"$A{row}",rows,cols),_diagram_formula("pm",f"$A{row}",rows,cols),_diagram_formula("am",f"$A{row}",rows,cols)) for row,code in enumerate(codes,2)]
    for ri,values in enumerate(content,1):
        row=doc.createElementNS(MAIN_NS,"row"); row.setAttribute("r",str(ri))
        for ci,value in enumerate(values,1):
            cell=doc.createElementNS(MAIN_NS,"c"); cell.setAttribute("r",f"{get_column_letter(ci)}{ri}")
            if ri==1: cell.setAttribute("s",str(header_style))
            _set_value(doc,cell,value,formula=isinstance(value,str) and value.startswith("=")); row.appendChild(cell)
        data.appendChild(row)
    doc.documentElement.appendChild(data)
    cells=_cellmap(doc); values=movement_values_by_code(frame,list(codes),metadata)
    for row,code in enumerate(codes,2):
        for column,key in (("B","total_12_hour"),("C","pm_peak"),("D","am_peak")):
            _set_formula_cache(doc,cells,f"{column}{row}",values[str(code)][key])
    return _xml(doc)


def _diagram_sheet_from_plan(diagram_rows,header_style):
    doc=minidom.getDOMImplementation().createDocument(MAIN_NS,"worksheet",None)
    _worksheet_shell(doc,4,len(diagram_rows)+1)
    data=doc.createElementNS(MAIN_NS,"sheetData")
    rows=[("movement_code","total_pcu","pm_peak_pcu","am_peak_pcu")]
    rows.extend((item.code,item.total_formula,item.pm_formula,item.am_formula) for item in diagram_rows)
    for ri,values in enumerate(rows,1):
        row=doc.createElementNS(MAIN_NS,"row"); row.setAttribute("r",str(ri))
        for ci,value in enumerate(values,1):
            cell=doc.createElementNS(MAIN_NS,"c"); cell.setAttribute("r",f"{get_column_letter(ci)}{ri}")
            if ri==1: cell.setAttribute("s",str(header_style))
            _set_value(doc,cell,value,formula=isinstance(value,str) and value.startswith("=")); row.appendChild(cell)
        data.appendChild(row)
    doc.documentElement.appendChild(data)
    cells=_cellmap(doc)
    for item in diagram_rows:
        for column,value in (("B",item.total_cache),("C",item.pm_cache),("D",item.am_cache)):
            _set_formula_cache(doc,cells,f"{column}{item.row}",value)
    return _xml(doc)

def _header_style(styles_xml):
    doc=minidom.parseString(styles_xml); root=doc.documentElement; fonts=_child(root,MAIN_NS,"fonts"); fills=_child(root,MAIN_NS,"fills"); xfs=_child(root,MAIN_NS,"cellXfs")
    if fonts is None or fills is None or xfs is None: raise ValueError("styles.xml missing required style collections")
    font_id=len(_els(fonts,MAIN_NS,"font")); fill_id=len(_els(fills,MAIN_NS,"fill")); style_id=len(_els(xfs,MAIN_NS,"xf"))
    font=_els(fonts,MAIN_NS,"font")[0].cloneNode(deep=True); _remove(font,MAIN_NS,"b","color"); bold=doc.createElementNS(MAIN_NS,"b"); color=doc.createElementNS(MAIN_NS,"color"); color.setAttribute("rgb","FFFFFFFF"); font.insertBefore(bold,font.firstChild); before=next((x for x in font.childNodes if x.nodeType==Node.ELEMENT_NODE and x.localName in {"sz","name","family","charset","scheme"}),None); font.insertBefore(color,before) if before else font.appendChild(color); fonts.appendChild(font); fonts.setAttribute("count",str(font_id+1))
    fill=doc.createElementNS(MAIN_NS,"fill"); pattern=doc.createElementNS(MAIN_NS,"patternFill"); pattern.setAttribute("patternType","solid"); fg=doc.createElementNS(MAIN_NS,"fgColor"); fg.setAttribute("rgb","FF1F4E78"); bg=doc.createElementNS(MAIN_NS,"bgColor"); bg.setAttribute("indexed","64"); pattern.appendChild(fg); pattern.appendChild(bg); fill.appendChild(pattern); fills.appendChild(fill); fills.setAttribute("count",str(fill_id+1))
    xf=doc.createElementNS(MAIN_NS,"xf")
    for key,value in (("numFmtId","0"),("fontId",str(font_id)),("fillId",str(fill_id)),("borderId","0"),("xfId","0"),("applyFont","1"),("applyFill","1")): xf.setAttribute(key,value)
    xfs.appendChild(xf); xfs.setAttribute("count",str(style_id+1)); return _xml(doc),style_id
def _chart_cache(chart_xml,categories,values):
    doc=minidom.parseString(chart_xml); series=_els(doc,CHART_NS,"ser")
    if len(series)!=1: raise ValueError(f"Expected one native chart series; found {len(series)}")
    for kind,data in (("cat",categories),("val",values)):
        box=_child(series[0],CHART_NS,kind); refs=_els(box,CHART_NS,"strRef" if kind=="cat" else "numRef") if box else []
        if len(refs)!=1: raise ValueError(f"Expected one chart {kind} range reference")
        cache_name="strCache" if kind=="cat" else "numCache"; caches=_els(refs[0],CHART_NS,cache_name); cache=caches[0] if caches else doc.createElementNS(CHART_NS,"c:"+cache_name)
        for x in list(cache.childNodes):
            if x.nodeType==Node.ELEMENT_NODE and x.namespaceURI==CHART_NS and x.localName in {"ptCount","pt"}: cache.removeChild(x); x.unlink()
        count=doc.createElementNS(CHART_NS,"c:ptCount"); count.setAttribute("val",str(len(data))); first=next((x for x in cache.childNodes if x.nodeType==Node.ELEMENT_NODE),None)
        cache.insertBefore(count,first) if first else cache.appendChild(count)
        for i,value in enumerate(data):
            p=doc.createElementNS(CHART_NS,"c:pt"); p.setAttribute("idx",str(i)); v=doc.createElementNS(CHART_NS,"c:v"); scalar=_scalar(value); v.appendChild(doc.createTextNode("" if scalar=="" else str(scalar))); p.appendChild(v); cache.appendChild(p)
        if not caches: refs[0].appendChild(cache)
    return _xml(doc)
def _app_props(app_xml,names):
    ep="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"; doc=minidom.parseString(app_xml); root=doc.documentElement
    heads=_els(root,ep,"HeadingPairs")
    if heads:
        vectors=_els(heads[0],VT_NS,"vector")
        if vectors:
            vs=_els(vectors[0],VT_NS,"variant")
            for i,v in enumerate(vs[:-1]):
                labels=_els(v,VT_NS,"lpstr")
                if labels and _text(labels[0])=="Worksheets":
                    counts=_els(vs[i+1],VT_NS,"i4")
                    if counts:
                        for child in list(counts[0].childNodes): counts[0].removeChild(child); child.unlink()
                        counts[0].appendChild(doc.createTextNode(str(len(names))))
                    break
    titles=_els(root,ep,"TitlesOfParts")
    if titles:
        vectors=_els(titles[0],VT_NS,"vector")
        if vectors:
            vec=vectors[0]
            for item in list(vec.childNodes):
                if item.nodeType==Node.ELEMENT_NODE and item.namespaceURI==VT_NS and item.localName=="lpstr": vec.removeChild(item); item.unlink()
            for name in names:
                entry=doc.createElementNS(VT_NS,"vt:lpstr"); entry.appendChild(doc.createTextNode(name)); vec.appendChild(entry)
            vec.setAttribute("size",str(len(names)))
    return _xml(doc)
def _structure_parts(z,add_names,sheet_parts):
    repl={}; added={}; removed={"xl/calcChain.xml"} if "xl/calcChain.xml" in z.namelist() else set()
    wb=minidom.parseString(z.read("xl/workbook.xml")); sheets=_els(wb,MAIN_NS,"sheets")[0]; ids=[int(s.getAttribute("sheetId")) for s in _els(sheets,MAIN_NS,"sheet")]; sheet_id=max(ids,default=0)+1
    rel=minidom.parseString(z.read("xl/_rels/workbook.xml.rels")); rels=_els(rel,PKG_REL_NS,"Relationship"); used={r.getAttribute("Id") for r in rels}; nums=[int(x[3:]) for x in used if x.startswith("rId") and x[3:].isdigit()]; ridnum=max(nums,default=0)+1
    partnums=[int(p.rsplit("sheet",1)[1].split(".",1)[0]) for p in sheet_parts.values() if "/sheet" in p and p.rsplit("sheet",1)[1].split(".",1)[0].isdigit()]; partnum=max(partnums,default=0)+1
    ct=minidom.parseString(z.read("[Content_Types].xml"))
    for name in add_names:
        part=f"xl/worksheets/sheet{partnum}.xml"
        while part in z.namelist() or part in added: partnum+=1; part=f"xl/worksheets/sheet{partnum}.xml"
        rid=f"rId{ridnum}"
        while rid in used: ridnum+=1; rid=f"rId{ridnum}"
        s=wb.createElementNS(MAIN_NS,"sheet"); s.setAttribute("name",name); s.setAttribute("sheetId",str(sheet_id)); s.setAttributeNS(DOC_REL_NS,"r:id",rid); sheets.appendChild(s)
        r=rel.createElementNS(PKG_REL_NS,"Relationship"); r.setAttribute("Id",rid); r.setAttribute("Type",WORKSHEET_REL); r.setAttribute("Target",f"worksheets/sheet{partnum}.xml"); rel.documentElement.appendChild(r); used.add(rid)
        o=ct.createElementNS(CONTENT_NS,"Override"); o.setAttribute("PartName","/"+part); o.setAttribute("ContentType",WORKSHEET_CONTENT_TYPE); ct.documentElement.appendChild(o)
        added[part]=b""; sheet_parts[name]=part; sheet_id+=1; ridnum+=1; partnum+=1
    for r in list(_els(rel,PKG_REL_NS,"Relationship")):
        if r.getAttribute("Type")==CALC_CHAIN_REL: rel.documentElement.removeChild(r); r.unlink()
    for o in list(_els(ct,CONTENT_NS,"Override")):
        if o.getAttribute("PartName")=="/xl/calcChain.xml": ct.documentElement.removeChild(o); o.unlink()
    calcs=_els(wb,MAIN_NS,"calcPr"); calc=calcs[0] if calcs else wb.createElementNS(MAIN_NS,"calcPr")
    if not calcs: wb.documentElement.appendChild(calc)
    for k in ("calcMode","fullCalcOnLoad","forceFullCalc"): calc.setAttribute(k,{"calcMode":"auto","fullCalcOnLoad":"1","forceFullCalc":"1"}[k])
    names=[s.getAttribute("name") for s in _els(sheets,MAIN_NS,"sheet")]
    repl.update({"xl/workbook.xml":_xml(wb),"xl/_rels/workbook.xml.rels":_xml(rel),"[Content_Types].xml":_xml(ct)})
    if add_names and "docProps/app.xml" in z.namelist(): repl["docProps/app.xml"]=_app_props(z.read("docProps/app.xml"),names)
    return repl,added,removed,names

def _write_metadata(doc,cells,mapping,metadata):
    values=_metadata_values(metadata); written=[]
    for key,info in mapping.get("metadata_cells",{}).items():
        ref=info.get("value_cell") or info.get("cell")
        if _put_summary(doc,cells,ref,values.get(key,""),mapping,intent=f"metadata.{key}",preserve=True): written.append(ref)
    move=mapping.get("movement_diagram_cells",{}); candidates=[(move.get("diagram_title",{}).get("cell"),values.get("survey_point",""),"movement.diagram_title"),(move.get("diagram_date",{}).get("cell"),values.get("survey_date",""),"movement.diagram_date")]
    for key,info in move.get("direction_labels",{}).items(): candidates.append((info.get("cell"),_mapped_label_value(metadata,key),f"movement.direction.{key}"))
    for key,info in move.get("road_labels",{}).items(): candidates.append((info.get("cell"),_mapped_label_value(metadata,key),f"movement.road.{key}"))
    candidates.append((move.get("caption",{}).get("cell"),metadata.get("caption_text") or "","movement.caption"))
    for ref,value,intent in candidates:
        if _put_summary(doc,cells,ref,value,mapping,intent=intent,preserve=True): written.append(ref)
    return written
def _write_support_payload(name,frame,header_style): return _dataframe_sheet(frame,header_style)
def export_template_ooxml(template_path,output_path,template_map,report_data,metadata,chart_source_data):
    """Patch the v1 template package directly; not routed from production yet."""
    source=Path(template_path).resolve(); target=Path(output_path).resolve()
    if source==target: raise ValueError("OOXML output must not overwrite the template.")
    if not source.is_file(): raise FileNotFoundError(source)
    if template_map.get("template_sheet","Summary")!="Summary": raise ValueError("Only the authoritative v1 Summary template is supported.")
    plan=resolve_template_write_plan(source,template_map,report_data,metadata,chart_source_data)
    target.parent.mkdir(parents=True,exist_ok=True); shutil.copy2(source,target)
    with zipfile.ZipFile(source) as original:
        names=set(original.namelist()); parts=_sheet_parts(original); summary_part=parts.get("Summary")
        if not summary_part: raise ValueError("Template has no Summary worksheet.")
        summary=minidom.parseString(original.read(summary_part)); cells=_cellmap(summary)
        formulas_before={c.getAttribute("r") for c in _els(summary,MAIN_NS,"c") if _child(c,MAIN_NS,"f") is not None}
        for write in plan.summary_writes:
            _put_summary(summary,cells,write.cell,write.value,template_map,intent=write.source,force=write.force,preserve=True)
        for binding in plan.formula_bindings:
            cell=cells.get(binding.cell)
            formula=_child(cell,MAIN_NS,"f") if cell is not None else None
            if formula is None:
                raise ValueError(f"Unknown native HLOOKUP formula at Summary!{binding.cell}.")
            if not _text(formula) and formula.getAttribute("t")=="shared":
                # The master formula is also in the 32-cell binding plan. Keep
                # the Excel-authored shared stub; its expanded formula follows it.
                continue
            if _text(formula)!=binding.old_formula[1:]:
                raise ValueError(f"Unknown native HLOOKUP formula at Summary!{binding.cell}.")
            for child in list(formula.childNodes): formula.removeChild(child); child.unlink()
            formula.appendChild(summary.createTextNode(binding.new_formula[1:]))
        preserved=tuple(sorted(f"Summary!{c.getAttribute('r')}" for c in _els(summary,MAIN_NS,"c") if _child(c,MAIN_NS,"f") is not None))
        formulas_after={c.getAttribute("r") for c in _els(summary,MAIN_NS,"c") if _child(c,MAIN_NS,"f") is not None}
        overwritten=tuple(sorted(f"Summary!{ref}" for ref in formulas_before-formulas_after if _in_range(ref,template_map.get("formula_overwrite_ranges",[]))))
        planned_chart_data={key:{"categories":list(categories),"values":list(values)} for key,categories,values in plan.chart_caches}
        _clear_formula_caches(summary); _cache_visible_report_formulas(summary,cells,_shared(original),template_map,planned_chart_data)
        movement_values={item.code:{"total_12_hour":item.total_cache,"pm_peak":item.pm_cache,"am_peak":item.am_cache} for item in plan.diagram_rows}
        _cache_summary_movement_values(summary,cells,_shared(original),template_map,plan,movement_values)
        _cache_required_template_sources(summary,cells,_shared(original),template_map,plan)
        replacements={summary_part:_xml(summary)}
        support=dict(plan.support_sheets); diagram_name=plan.diagram_sheet_name
        if "Summary" in support or diagram_name in support: raise ValueError("Support sheet names collide with protected worksheet names.")
        collisions=set(support)|{diagram_name}; existing=set(parts)
        if collisions & existing: raise ValueError(f"Refusing to replace existing template sheets: {sorted(collisions & existing)}")
        add_names=list(support)+[diagram_name]
        structure,added,removed,sheet_names=_structure_parts(original,add_names,parts); replacements.update(structure)
        styles,header_style=_header_style(original.read("xl/styles.xml")); replacements["xl/styles.xml"]=styles
        for name,frame in support.items(): added[parts[name]]=_write_support_payload(name,frame,header_style)
        added[parts[diagram_name]]=_diagram_sheet_from_plan(plan.diagram_rows,header_style)
        chart_parts={}
        anchors=template_map.get("chart_anchors",{})
        for anchor,data_key in (("hourly_pcu_line_chart","hourly_pcu"),("vehicle_composition_bar_chart","vehicle_composition")):
            part=(anchors.get(anchor) or {}).get("chart_xml")
            if part not in names: raise ValueError(f"Mapped native chart part missing: {part}")
            values={key:(categories,series) for key,categories,series in plan.chart_caches}[data_key]
            chart_parts[part]=_chart_cache(original.read(part),list(values[0]),list(values[1]))
        replacements.update(chart_parts)
        with zipfile.ZipFile(target,"w") as output:
            output.comment=original.comment
            for info in original.infolist():
                if info.filename in removed: continue
                output.writestr(info,replacements.get(info.filename,original.read(info.filename)))
            for part in sorted(added,key=lambda p:int(p.rsplit("sheet",1)[1].split(".",1)[0])):
                info=zipfile.ZipInfo(part,date_time=(1980,1,1,0,0,0)); info.compress_type=zipfile.ZIP_DEFLATED; output.writestr(info,added[part])
    with zipfile.ZipFile(target) as result:
        bad=result.testzip()
        if bad: raise ValueError(f"ZIP CRC error: {bad}")
        actual=_sheet_parts(result)
        if list(actual)!=sheet_names: raise ValueError("Allocated workbook sheets differ from resulting relationships.")
        wb=minidom.parseString(result.read("xl/workbook.xml")); calc=_els(wb,MAIN_NS,"calcPr")[0]
        if any(calc.getAttribute(k)!=v for k,v in (("calcMode","auto"),("fullCalcOnLoad","1"),("forceFullCalc","1"))): raise ValueError("Workbook recalc flags are incomplete.")
        if "xl/calcChain.xml" in result.namelist(): raise ValueError("Stale calcChain remains in generated workbook.")
    with zipfile.ZipFile(source) as original:
        changed=tuple(sorted(p for p in replacements if p in original.namelist()))
    return OOXMLExportResult(target,changed,tuple(sorted(added)),tuple(sorted(removed)),preserved,overwritten)
