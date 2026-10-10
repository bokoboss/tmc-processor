"""Independent package-level checks for a resolved native-template write plan."""

from __future__ import annotations

import math
import posixpath
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from xml.dom import minidom

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

from .ooxml_template_export import (
    CHART_NS,
    MAIN_NS,
    PKG_REL_NS,
    _cellmap,
    _cellvalue,
    _chart_cache,
    _child,
    _date_serial,
    _els,
    _scalar,
    _shared,
    _sheet_parts,
    _text,
    _range_refs,
)
from .template_write_plan import TemplateWritePlan, v2_final_total_formulas


def _expected(value):
    value = _scalar(value)
    if isinstance(value, datetime):
        return _date_serial(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return value


def _equal(actual, expected):
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12)
    return actual == expected


def _relationship_target(rels_path: str, target: str) -> str:
    parent = posixpath.dirname(rels_path)
    if parent.endswith("/_rels"):
        parent = parent[:-6]
    elif parent == "_rels":
        parent = ""
    return posixpath.normpath(posixpath.join(parent, target)) if not target.startswith("/") else target.lstrip("/")


def verify_ooxml_against_plan(
    template_path: str | Path,
    output_path: str | Path,
    template_map: dict,
    plan: TemplateWritePlan,
) -> tuple[str, ...]:
    """Raise for a missing, misplaced, stale, or unexpected workbook write."""
    issues: list[str] = []
    with zipfile.ZipFile(template_path) as original, zipfile.ZipFile(output_path) as result:
        if result.testzip() is not None:
            issues.append("ZIP CRC failed")
        source_names, result_names = set(original.namelist()), set(result.namelist())
        allowed_changed = {
            "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml",
            "docProps/app.xml", "xl/styles.xml",
            _sheet_parts(original)[plan.template_sheet],
            template_map["chart_anchors"]["hourly_pcu_line_chart"]["chart_xml"],
            template_map["chart_anchors"]["vehicle_composition_bar_chart"]["chart_xml"],
        }
        unexpected = sorted(part for part in source_names & result_names if original.read(part) != result.read(part) and part not in allowed_changed)
        if unexpected:
            issues.append(f"Unexpected package changes: {unexpected}")
        unexpected_removed = sorted((source_names - result_names) - {"xl/calcChain.xml"})
        if unexpected_removed:
            issues.append(f"Unexpected removed parts: {unexpected_removed}")
        unexpected_added = sorted(part for part in result_names - source_names if not (part.startswith("xl/worksheets/sheet") and part.endswith(".xml")))
        if unexpected_added:
            issues.append(f"Unexpected added parts: {unexpected_added}")
        for part in result_names:
            if part.endswith((".xml", ".rels")) or part == "[Content_Types].xml":
                try:
                    minidom.parseString(result.read(part))
                except Exception as exc:
                    issues.append(f"Invalid XML {part}: {exc}")
            if part.endswith(".rels"):
                root = minidom.parseString(result.read(part))
                for relation in _els(root, PKG_REL_NS, "Relationship"):
                    if relation.getAttribute("TargetMode").casefold() == "external":
                        continue
                    target = _relationship_target(part, relation.getAttribute("Target"))
                    if target not in result_names:
                        issues.append(f"Broken relationship {part} -> {target}")
        try:
            parts = _sheet_parts(result)
        except ValueError as exc:
            raise AssertionError(f"OOXML write-plan verification failed:\n{exc}") from exc
        template_parts = _sheet_parts(original)
        expected_names = [plan.template_sheet, *(name for name, _ in plan.support_sheets), plan.diagram_sheet_name]
        if list(parts) != expected_names:
            issues.append(f"Sheet order/content differs: {list(parts)}")
        if plan.template_sheet not in parts:
            raise AssertionError("Summary sheet missing")
        strings = _shared(result)
        summary = _cellmap(minidom.parseString(result.read(parts[plan.template_sheet])))
        original_summary = _cellmap(minidom.parseString(original.read(template_parts[plan.template_sheet])))
        final_writes = {write.cell: write for write in plan.summary_writes}
        final_totals = {item.cell: item for item in plan.final_total_formulas}
        original_formulas = {ref: "=" + _text(_child(cell, MAIN_NS, "f")) for ref, cell in original_summary.items() if _child(cell, MAIN_NS, "f") is not None}
        expected_totals = v2_final_total_formulas(original_formulas, dict(plan.support_sheets)) if template_map.get("movement_code_scheme") == "approach_movement" else ()
        if plan.final_total_formulas != expected_totals:
            issues.append("Final PCU overrides differ from the validated V2 allow-list")
        for item in expected_totals:
            formula = _child(summary.get(item.cell), MAIN_NS, "f")
            if formula is None or _text(formula) != item.new_formula[1:] or not _equal(_cellvalue(summary.get(item.cell), strings), item.cache):
                issues.append(f"{plan.template_sheet}!{item.cell}: final PCU formula/cache mismatch")
        for ref, expanded in plan.preserved_formulas:
            before = _child(original_summary.get(ref), MAIN_NS, "f")
            after = _child(summary.get(ref), MAIN_NS, "f")
            if "W40" in final_totals and ref in _range_refs("X40:AJ40"):
                if before is None or dict(before.attributes.items()) != {"t": "shared", "si": "8"} or _text(before) or after is None or after.attributes.length or _text(after) != expanded[1:]:
                    issues.append(f"{plan.template_sheet}!{ref}: shared vehicle total materialization changed formula")
                continue
            if before is None or after is None or before.toxml() != after.toxml():
                issues.append(f"{plan.template_sheet}!{ref}: protected formula XML changed")
        for row in range(9, 23):
            ref = f"U{row}"
            if _cellvalue(original_summary.get(ref), _shared(original)) != row - 8 or _cellvalue(summary.get(ref), strings) != row - 8:
                issues.append(f"{plan.template_sheet}!{ref}: immutable helper sequence must be 1..14")
        peak_rows = {peak.period: peak for peak in plan.peak_formula_rows}
        if len(plan.formula_bindings) != 32 or len({binding.cell for binding in plan.formula_bindings}) != 32:
            issues.append("Native Peak binding plan must contain 32 unique AM/PM formulas")
        for binding in plan.formula_bindings:
            before = _child(original_summary.get(binding.cell), MAIN_NS, "f")
            after = _child(summary.get(binding.cell), MAIN_NS, "f")
            if before is None or after is None or before.toxml().replace(binding.old_formula[1:], binding.new_formula[1:]) != after.toxml():
                issues.append(f"{plan.template_sheet}!{binding.cell}: native Peak HLOOKUP binding changed unexpectedly")
            peak = peak_rows[binding.period]
            column = template_map["hourly_movement_table"]["columns"][binding.code]
            selected = _cellvalue(summary.get(f"{column}{peak.worksheet_row}"), strings)
            actual = _cellvalue(summary.get(binding.cell), strings)
            if not _equal(actual, selected):
                issues.append(f"{plan.template_sheet}!{binding.cell}: Peak cache differs from selected row {peak.worksheet_row}")
        for ref, write in final_writes.items():
            actual = _cellvalue(summary.get(ref), strings)
            expected = _expected(write.value)
            if not _equal(actual, expected):
                issues.append(f"{plan.template_sheet}!{ref}: {write.source}: {actual!r} != {expected!r}")
        diagram_values = {item.code: item for item in plan.diagram_rows}
        for code, mapping in template_map["movement_diagram_cells"]["diagram_movements"].items():
            item = diagram_values[code]
            for key, value in (("total_12_hour_cell", item.total_cache), ("pm_peak_hour_cell", item.pm_cache), ("am_peak_hour_cell", item.am_cache)):
                ref = mapping.get(key)
                if ref and not _equal(_cellvalue(summary.get(ref), strings), value):
                    issues.append(f"{plan.template_sheet}!{ref}: stale movement formula cache")
        movement_table = template_map["hourly_movement_table"]
        vehicle_table = template_map["hourly_vehicle_class_table"]
        vehicle_first = int(vehicle_table["first_data_row"])
        vehicle_last = int(vehicle_table["last_data_row"])
        vehicle_pcu = vehicle_table["columns"]["total_pcu"]
        movement_total = movement_table["columns"]["Total"]
        movement_first = int(movement_table["first_data_row"])
        movement_last = int(movement_table["last_data_row"])
        movement_total_row = int(movement_table["total_row"])
        diagram_by_code = {item.code: item for item in plan.diagram_rows}
        hourly_support = dict(plan.support_sheets).get("Hourly_Movement_PCU")
        if hourly_support is None:
            issues.append("Hourly_Movement_PCU support data missing from write plan")
        for code, column in movement_table["columns"].items():
            if code == "time":
                continue
            ref = f"{column}{movement_total_row}"
            if ref in final_totals:
                continue
            values = [_cellvalue(summary.get(f"{column}{row}"), strings) for row in range(movement_first, movement_last + 1)]
            if not all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values):
                issues.append(f"{plan.template_sheet}!{ref}: non-numeric hourly movement source")
                continue
            expected = sum(values)
            if code != "Total":
                item = diagram_by_code.get(code)
                if item is None or not _equal(expected, item.total_cache):
                    issues.append(f"{plan.template_sheet}!{ref}: movement total disagrees with Diagram_Data for {code}")
            if hourly_support is not None and code in hourly_support.columns:
                support_total = hourly_support.iloc[-1][code]
                # V2's authored cells display integer PCU values; summing its 12
                # hourly display values can differ from the precise support total
                # by at most half a PCU per hour.
                rounded_v2_total = (
                    code == "Total"
                    and template_map.get("movement_code_scheme") == "approach_movement"
                    and math.isclose(expected, float(support_total), rel_tol=1e-9, abs_tol=6.0)
                )
                if not _equal(expected, support_total) and not rounded_v2_total:
                    issues.append(f"{plan.template_sheet}!{ref}: total disagrees with Hourly_Movement_PCU support data")
            cell = summary.get(ref)
            formula = _child(cell, MAIN_NS, "f") if cell is not None else None
            if formula is None or (_text(formula) and _text(formula) != f"SUM({column}{movement_first}:{column}{movement_last})") or (not _text(formula) and formula.getAttribute("t") != "shared"):
                issues.append(f"{plan.template_sheet}!{ref}: native total formula changed")
            actual = _cellvalue(cell, strings)
            if actual is None:
                issues.append(f"{plan.template_sheet}!{ref}: required formula cache missing")
            elif not _equal(actual, expected):
                issues.append(f"{plan.template_sheet}!{ref}: wrong formula cache {actual!r} != {expected!r}")

        rollups = {
            "F30": ("G19:G22", "J16:M16", "O21:O24", "J27:M27"),
            "F31": ("J15:M15", "F19:F22", "J28:M28", "P21:P24"),
            "F32": ("J14:M14", "E19:E22", "J29:M29", "Q21:Q24"),
        }
        for ref, areas in rollups.items():
            if ref in final_totals:
                continue
            sources = [_cellvalue(summary.get(source), strings) for area in areas for source in _range_refs(area)]
            if not all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in sources):
                issues.append(f"{plan.template_sheet}!{ref}: movement rollup source cache missing")
                continue
            expected = sum(sources)
            cell = summary.get(ref)
            formula = _child(cell, MAIN_NS, "f") if cell is not None else None
            if formula is None or _text(formula) != f"SUM({','.join(areas)})":
                issues.append(f"{plan.template_sheet}!{ref}: native rollup formula changed")
            actual = _cellvalue(cell, strings)
            if actual is None:
                issues.append(f"{plan.template_sheet}!{ref}: required formula cache missing")
            elif not _equal(actual, expected):
                issues.append(f"{plan.template_sheet}!{ref}: wrong formula cache {actual!r} != {expected!r}")
        title_cache = _cellvalue(summary.get("C8"), strings)
        title_source = _cellvalue(summary.get("E5"), strings) or ""
        if title_cache is None or title_cache != title_source:
            issues.append(f"{plan.template_sheet}!C8: diagram title cache missing or stale")
        date_source = _cellvalue(summary.get("K5"), strings)
        if date_source is None or isinstance(date_source, (int, float)) and not isinstance(date_source, bool):
            expected = 0 if date_source is None else date_source
            if not _equal(_cellvalue(summary.get("C9"), strings), expected):
                issues.append(f"{plan.template_sheet}!C9: numeric diagram date cache missing or stale")
        else:
            date_cell = summary.get("C9")
            date_cache = _child(date_cell, MAIN_NS, "v") if date_cell is not None else None
            if date_cache is not None and _text(date_cache):
                issues.append(f"{plan.template_sheet}!C9: text date has an unverified formula cache")

        excluded_date = isinstance(date_source, str)
        for ref, cell in summary.items():
            if _child(cell, MAIN_NS, "f") is None or excluded_date and ref == "C9":
                continue
            cached = _child(cell, MAIN_NS, "v")
            if cached is None or not _text(cached) and cell.getAttribute("t") != "str":
                issues.append(f"{plan.template_sheet}!{ref}: required formula cache missing")

        for offset, row in enumerate(range(vehicle_first, vehicle_last + 1)):
            expected = _cellvalue(summary.get(f"{movement_total}{movement_first + offset}"), strings)
            if not _equal(_cellvalue(summary.get(f"{vehicle_pcu}{row}"), strings), expected):
                issues.append(f"{plan.template_sheet}!{vehicle_pcu}{row}: stale hourly PCU cache")
        for key, column in vehicle_table["columns"].items():
            if key == "time":
                continue
            expected = sum(float(_cellvalue(summary.get(f"{column}{row}"), strings) or 0) for row in range(vehicle_first, vehicle_last + 1))
            ref = f"{column}{vehicle_table['total_row']}"
            if ref in final_totals:
                continue
            if not _equal(_cellvalue(summary.get(ref), strings), expected):
                issues.append(f"{plan.template_sheet}!{ref}: stale vehicle total cache")
        for name, frame in plan.support_sheets:
            if name not in parts:
                issues.append(f"Missing support sheet {name}")
                continue
            cells = _cellmap(minidom.parseString(result.read(parts[name])))
            rows = [tuple(map(str, frame.columns)), *frame.itertuples(index=False, name=None)]
            if not len(frame.columns):
                rows = [("",)]
            for row_index, row in enumerate(rows, 1):
                for column_index, value in enumerate(row, 1):
                    ref = f"{get_column_letter(column_index)}{row_index}"
                    if not _equal(_cellvalue(cells.get(ref), strings), _expected(value)):
                        issues.append(f"{name}!{ref}: support value mismatch")
        if plan.diagram_sheet_name in parts:
            diagram = _cellmap(minidom.parseString(result.read(parts[plan.diagram_sheet_name])))
            for item in plan.diagram_rows:
                if _cellvalue(diagram.get(f"A{item.row}"), strings) != item.code:
                    issues.append(f"{plan.diagram_sheet_name}!A{item.row}: movement code mismatch")
                for column, formula, cache in (("B", item.total_formula, item.total_cache), ("C", item.pm_formula, item.pm_cache), ("D", item.am_formula, item.am_cache)):
                    ref = f"{column}{item.row}"
                    cell = diagram.get(ref)
                    if cell is None or "=" + _text(_child(cell, MAIN_NS, "f")) != formula:
                        issues.append(f"{plan.diagram_sheet_name}!{ref}: formula mismatch")
                    if not _equal(_cellvalue(cell, strings), cache):
                        issues.append(f"{plan.diagram_sheet_name}!{ref}: stale formula cache")
        anchors = template_map["chart_anchors"]
        cache_data = {key: (categories, values) for key, categories, values in plan.chart_caches}
        vehicle_source = anchors["native_vehicle_composition_chart_source"]
        for category_ref in _range_refs(vehicle_source["categories"]):
            column = "".join(char for char in category_ref if char.isalpha())
            row = int("".join(char for char in category_ref if char.isdigit()))
            header = _cellvalue(summary.get(category_ref), strings)
            vehicle_columns = vehicle_table["columns"]
            matched = next((source_col for key, source_col in vehicle_columns.items() if key not in {"time", "total_pcu", "total_vehicles"} and _cellvalue(summary.get(f"{source_col}{vehicle_table['header_row']}"), strings) == header), None)
            if matched is None or not _equal(_cellvalue(summary.get(f"{column}{row - 1}"), strings), _cellvalue(summary.get(f"{matched}{vehicle_table['total_row']}"), strings)):
                issues.append(f"{plan.template_sheet}!{column}{row - 1}: stale vehicle chart numerator cache")
        vehicle_values = cache_data["vehicle_composition"][1]
        for ref, value in zip(_range_refs(vehicle_source["values"]), vehicle_values):
            if not _equal(_cellvalue(summary.get(ref), strings), _expected(value)):
                issues.append(f"{plan.template_sheet}!{ref}: stale vehicle chart source cache")
        for anchor, key in (("hourly_pcu_line_chart", "hourly_pcu"), ("vehicle_composition_bar_chart", "vehicle_composition")):
            part = anchors[anchor]["chart_xml"]
            categories, values = cache_data[key]
            if part not in result_names or result.read(part) != _chart_cache(original.read(part), list(categories), list(values)):
                issues.append(f"{part}: chart structure or cache mismatch")
        for part in source_names:
            if part.startswith(("xl/drawings/", "xl/media/")) and (part not in result_names or original.read(part) != result.read(part)):
                issues.append(f"{part}: protected drawing/media changed")
        # Formula comparison uses openpyxl's shared-formula expansion, then
        # compares only values. Neither workbook is saved through openpyxl.
        source_book = load_workbook(template_path, read_only=True, data_only=False)
        output_book = load_workbook(output_path, read_only=True, data_only=False)
        try:
            source_sheet = source_book[plan.template_sheet]
            output_sheet = output_book[plan.template_sheet]
            for ref, formula in plan.preserved_formulas:
                if source_sheet[ref].value != formula or output_sheet[ref].value != formula:
                    issues.append(f"{plan.template_sheet}!{ref}: protected formula changed")
            for binding in plan.formula_bindings:
                if source_sheet[binding.cell].value != binding.old_formula or output_sheet[binding.cell].value != binding.new_formula:
                    issues.append(f"{plan.template_sheet}!{binding.cell}: Peak formula differs from binding plan")
        finally:
            source_book.close()
            output_book.close()
    if issues:
        raise AssertionError("OOXML write-plan verification failed:\n" + "\n".join(issues[:40]))
    return ("planned_summary_writes", "support_sheets", "diagram_formulas_and_caches", "chart_caches", "protected_formulas", "drawings_media", "xml_relationships")
