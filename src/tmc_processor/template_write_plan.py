"""Resolve native-template writes before choosing an Excel transport.

The plan starts at the post-Review export payload. It does not derive traffic
movements or change Mapping, Analyze, or Review decisions.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, ROUND_HALF_UP
import math
from numbers import Number
from pathlib import Path
import re
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from .diagram import MOVEMENT_CODES
from .movement_scheme import APPROACH_MOVEMENT_CODES
from .report_template import (
    _column_for_key,
    _mapped_label_value,
    _metadata_values,
    _normalise_time_label,
    movement_values_by_code,
)
from .time_utils import time_to_minutes


@dataclass(frozen=True)
class CellWrite:
    sheet: str
    cell: str
    value: Any
    source: str
    order: int
    force: bool = False

    @property
    def value_type(self) -> str:
        value = self.value
        if value is None or value is pd.NA or value is pd.NaT:
            return "blank"
        try:
            if pd.isna(value):
                return "blank"
        except (TypeError, ValueError):
            pass
        if isinstance(value, bool):
            return "boolean"
        if isinstance(value, Number):
            return "number"
        if isinstance(value, (date, datetime, time, pd.Timestamp)):
            return "date_time"
        return "formula" if isinstance(value, str) and value.startswith("=") else "text"


@dataclass(frozen=True)
class DiagramRow:
    row: int
    code: str
    total_formula: str
    pm_formula: str
    am_formula: str
    total_cache: float
    pm_cache: float
    am_cache: float


@dataclass(frozen=True)
class PeakFormulaRow:
    period: str
    start: str
    end: str
    worksheet_row: int
    helper_cell: str
    helper_value: int


@dataclass(frozen=True)
class PeakBindingMatch:
    period: str
    interval: str
    matched: bool
    worksheet_row: int | None
    helper_cell: str | None
    helper_value: int | None


def _match_peak_rows(
    time_labels: dict[int, Any],
    hourly_map: dict[str, Any],
    helper_column: str,
    metadata: dict[str, Any],
) -> tuple[PeakBindingMatch, ...]:
    first, last = int(hourly_map["first_data_row"]), int(hourly_map["last_data_row"])
    matches: list[PeakBindingMatch] = []
    for period in ("am", "pm"):
        start = _normalise_time_label(metadata.get(f"{period}_peak_start"))
        end = _normalise_time_label(metadata.get(f"{period}_peak_end"))
        interval = f"{start}-{end}" if start and end else ""
        rows = [
            row for row in range(first, last + 1)
            if interval and _normalise_time_label(time_labels[row]) == interval
        ]
        row = rows[0] if len(rows) == 1 else None
        matches.append(PeakBindingMatch(
            period, interval, row is not None, row,
            f"{helper_column}{row}" if row is not None else None,
            row - int(hourly_map["header_row"]) + 1 if row is not None else None,
        ))
    return tuple(matches)


def preflight_template_peak_binding(
    template_path: str | Path,
    template_map: dict[str, Any],
    metadata: dict[str, Any],
) -> tuple[PeakBindingMatch, ...]:
    """Read the native Summary time rows without changing the workbook."""
    hourly_map = template_map["hourly_movement_table"]
    helper_column = str(template_map["movement_diagram_cells"]["movement_value_rows"].get("helper_column") or "U")
    workbook = load_workbook(template_path, read_only=True, data_only=False)
    try:
        summary = workbook[str(template_map.get("template_sheet") or "Summary")]
        if helper_column != "U" or [summary[f"{helper_column}{row}"].value for row in range(9, 23)] != list(range(1, 15)):
            raise ValueError("Authoritative Summary!U9:U22 helper sequence must be 1..14.")
        labels = {
            row: summary[f"{hourly_map['time_column']}{row}"].value
            for row in range(int(hourly_map["first_data_row"]), int(hourly_map["last_data_row"]) + 1)
        }
        return _match_peak_rows(labels, hourly_map, helper_column, metadata)
    finally:
        workbook.close()


@dataclass(frozen=True)
class FormulaBinding:
    period: str
    code: str
    cell: str
    old_formula: str
    new_formula: str


@dataclass(frozen=True)
class FinalTotalFormula:
    cell: str
    old_formula: str
    new_formula: str
    cache: int


def v2_final_total_formulas(formulas, support) -> tuple[FinalTotalFormula, ...]:
    """Allow only the five authored final PCU rollups to round precise sources."""
    movement = support["Movement_Summary"]
    if list(movement.columns) != ["movement_code_scheme", "movement_code", "approach_direction", "movement_type", "display_label", "count", "pcu"] or movement["movement_code"].tolist() != list(APPROACH_MOVEMENT_CODES):
        raise ValueError("V2 final totals require the canonical 16-row Movement_Summary contract.")
    amount = math.fsum(float(value) for value in movement["pcu"])
    for name in ("Normalized_Data", "Hourly_Totals"):
        precise = math.fsum(float(value) for value in support[name]["pcu"])
        if not math.isfinite(precise) or not math.isclose(amount, precise, rel_tol=1e-12, abs_tol=1e-9):
            raise ValueError(f"V2 precise Movement_Summary total disagrees with {name}.")
    peak = support["Peak_PHF"]
    if len(peak) != 1 or list(peak.columns)[5] != "am_peak_pcu" or list(peak.columns)[9] != "pm_peak_pcu":
        raise ValueError("V2 final totals require the precise Peak_PHF contract.")
    total_formula = "=ROUND(SUM('Movement_Summary'!$G$2:$G$17),0)"
    overrides = (
        ("F30", "=SUM(G19:G22,J16:M16,O21:O24,J27:M27)", "=ROUND('Peak_PHF'!$F$2,0)", peak.iloc[0]["am_peak_pcu"]),
        ("F31", "=SUM(J15:M15,F19:F22,J28:M28,P21:P24)", "=ROUND('Peak_PHF'!$J$2,0)", peak.iloc[0]["pm_peak_pcu"]),
        ("F32", "=SUM(J14:M14,E19:E22,J29:M29,Q21:Q24)", total_formula, amount),
        ("W40", "=SUM(W28:W39)", total_formula, amount),
        ("AM22", "=SUM(AM10:AM21)", total_formula, amount),
    )
    result = []
    for cell, old, new, precise in overrides:
        if formulas.get(cell) != old or not math.isfinite(float(precise)):
            raise ValueError(f"Unknown V2 final total formula/source at Summary!{cell}.")
        cache = int(Decimal(str(precise)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        result.append(FinalTotalFormula(cell, old, new, cache))
    return tuple(result)


def v1_final_total_formulas(formulas, support) -> tuple[FinalTotalFormula, ...]:
    """Validate V1's variable support rows before rounding the five final cells."""
    movement = support.get("Movement_Summary")
    normalized = support.get("Normalized_Data")
    hourly = support.get("Hourly_Summary")
    peak = support.get("Peak_PHF")
    if movement is None or list(movement.columns) != [
        "movement_code", "from_leg", "to_leg", "turn_type", "facility_type",
        "include_in_peak", "include_in_report", "count", "pcu",
    ] or movement.empty:
        raise ValueError("V1 final totals require the exported Movement_Summary schema and at least one data row.")
    if normalized is None or not {"movement_code", "output_movement_code", "from_leg", "to_leg", "turn_type", "facility_type", "include_in_peak", "include_in_report", "time_start", "pcu"}.issubset(normalized.columns):
        raise ValueError("V1 final totals require the exported Normalized_Data source contract.")
    if hourly is None or list(hourly.columns) != ["hour_start", "hour_end", "count", "pcu"]:
        raise ValueError("V1 final totals require the exported Hourly_Summary schema.")
    if peak is None or list(peak.columns) != [
        "peak_mode", "am_peak_search_window", "pm_peak_search_window", "am_peak_start",
        "am_peak_end", "am_peak_pcu", "am_phf", "pm_peak_start", "pm_peak_end",
        "pm_peak_pcu", "pm_phf", "peak_selection_source",
    ] or len(peak) != 1:
        raise ValueError("V1 final totals require the one-row Peak_PHF source contract.")

    def finite_values(frame, column, name):
        values = pd.to_numeric(frame[column], errors="coerce").to_numpy(dtype=float)
        if not len(values) or not all(math.isfinite(value) for value in values):
            raise ValueError(f"V1 final totals require finite {name} PCU values.")
        return values

    for frame, name in ((movement, "Movement_Summary"), (normalized, "Normalized_Data"), (hourly, "Hourly_Summary")):
        finite_values(frame, "pcu", name)
    if not movement["include_in_report"].map(lambda value: isinstance(value, (bool,)) or type(value).__name__ == "bool_").all():
        raise ValueError("V1 Movement_Summary include_in_report must contain booleans.")
    if not normalized["include_in_report"].map(lambda value: isinstance(value, (bool,)) or type(value).__name__ == "bool_").all():
        raise ValueError("V1 Normalized_Data include_in_report must contain booleans.")
    if not normalized["include_in_peak"].map(lambda value: isinstance(value, (bool,)) or type(value).__name__ == "bool_").all():
        raise ValueError("V1 Normalized_Data include_in_peak must contain booleans.")

    # Rebuild V1's movement grouping and compare every row, including excluded
    # report groups; this prevents a stale, duplicated, or missing support row
    # from silently changing the validated dynamic SUMIFS range.
    group_keys = ["movement_code", "from_leg", "to_leg", "turn_type", "facility_type", "include_in_peak", "include_in_report"]
    if movement.duplicated(subset=group_keys, keep=False).any():
        raise ValueError("V1 Movement_Summary contains duplicate canonical group keys.")
    grouped_source = normalized.copy()
    output_code = grouped_source["output_movement_code"].fillna("").astype(str).str.strip()
    source_code = grouped_source["movement_code"].fillna("").astype(str).str.strip()
    grouped_source["movement_code"] = output_code.where(output_code != "", source_code)
    expected_rows = grouped_source.groupby(group_keys, dropna=False, as_index=False).agg(pcu=("pcu", "sum"))
    actual_rows = movement[group_keys + ["pcu"]].groupby(group_keys, dropna=False, as_index=False).agg(pcu=("pcu", "sum"))
    if len(actual_rows) != len(expected_rows):
        raise ValueError("V1 Movement_Summary rows disagree with the normalized source grouping.")
    merged = expected_rows.merge(actual_rows, on=group_keys, how="outer", suffixes=("_source", "_summary"), indicator=True)
    if not (merged["_merge"].eq("both").all() and all(
        math.isclose(float(left), float(right), rel_tol=1e-12, abs_tol=1e-9)
        for left, right in zip(merged["pcu_source"], merged["pcu_summary"])
    )):
        raise ValueError("V1 Movement_Summary PCU rows disagree with Normalized_Data.")

    report_total = math.fsum(float(value) for value in movement.loc[movement["include_in_report"], "pcu"])
    normalized_report_total = math.fsum(float(value) for value in normalized.loc[normalized["include_in_report"], "pcu"])
    hourly_total = math.fsum(float(value) for value in hourly["pcu"])
    normalized_total = math.fsum(float(value) for value in normalized["pcu"])
    if not math.isclose(report_total, normalized_report_total, rel_tol=1e-12, abs_tol=1e-9):
        raise ValueError("V1 reported Movement_Summary PCU disagrees with Normalized_Data.")
    if not math.isclose(hourly_total, normalized_total, rel_tol=1e-12, abs_tol=1e-9):
        raise ValueError("V1 Hourly_Summary PCU disagrees with Normalized_Data.")

    normalized_by_hour = {}
    for time_start, pcu in zip(normalized["time_start"], normalized["pcu"]):
        minute = time_to_minutes(time_start)
        if minute is None:
            raise ValueError("V1 Normalized_Data contains an invalid source hour.")
        hour_start = (minute // 60) * 60
        normalized_by_hour.setdefault(hour_start, []).append(float(pcu))
    normalized_by_hour = {hour: math.fsum(values) for hour, values in normalized_by_hour.items()}

    summary_by_hour = {}
    for row in hourly.itertuples(index=False):
        start = time_to_minutes(row.hour_start)
        end = time_to_minutes(row.hour_end)
        if start is None or end is None or start % 60 != 0 or end != (start + 60) % (24 * 60):
            raise ValueError("V1 Hourly_Summary contains an invalid hourly interval.")
        if start in summary_by_hour:
            raise ValueError("V1 Hourly_Summary contains duplicate hourly intervals.")
        summary_by_hour[start] = float(row.pcu)
    # Compare every source and summary hour independently. The 1e-12 relative
    # and 1e-9 absolute tolerance accommodates only binary floating summation
    # noise; it cannot hide a rounded or redistributed PCU contribution.
    for hour in normalized_by_hour.keys() | summary_by_hour.keys():
        precise_source = normalized_by_hour.get(hour, 0.0)
        precise_summary = summary_by_hour.get(hour, 0.0)
        if not math.isclose(precise_source, precise_summary, rel_tol=1e-12, abs_tol=1e-9):
            raise ValueError(
                f"V1 Hourly_Summary PCU disagrees with Normalized_Data for source hour {hour:04d}."
            )

    if not all(math.isfinite(float(peak.iloc[0][name])) for name in ("am_peak_pcu", "pm_peak_pcu")):
        raise ValueError("V1 final totals require finite effective Peak_PHF PCU values.")
    for period in ("am", "pm"):
        start = time_to_minutes(peak.iloc[0][f"{period}_peak_start"])
        end = time_to_minutes(peak.iloc[0][f"{period}_peak_end"])
        if start is None or end is None or end - start != 60:
            raise ValueError(f"V1 {period.upper()} Peak_PHF must identify one effective 60-minute interval.")
        source = normalized.loc[normalized["include_in_peak"]].copy()
        source["_minute"] = source["time_start"].map(time_to_minutes)
        precise_peak = math.fsum(float(value) for value in source.loc[source["_minute"].between(start, end - 1), "pcu"])
        if not math.isclose(precise_peak, float(peak.iloc[0][f"{period}_peak_pcu"]), rel_tol=1e-12, abs_tol=1e-9):
            raise ValueError(f"V1 effective {period.upper()} Peak_PHF disagrees with normalized include_in_peak PCU.")
    total_formula = "=ROUND(SUMIFS('Movement_Summary'!$I$2:$I${last},'Movement_Summary'!$G$2:$G${last},TRUE),0)"
    result = []
    last = len(movement) + 1
    total_formula = total_formula.format(last=last)
    overrides = (
        ("F30", "=SUM(G19:G22,J16:M16,O21:O24,J27:M27)", "=ROUND('Peak_PHF'!$F$2,0)", peak.iloc[0]["am_peak_pcu"]),
        ("F31", "=SUM(J15:M15,F19:F22,J28:M28,P21:P24)", "=ROUND('Peak_PHF'!$J$2,0)", peak.iloc[0]["pm_peak_pcu"]),
        ("F32", "=SUM(J14:M14,E19:E22,J29:M29,Q21:Q24)", total_formula, report_total),
        ("W40", "=SUM(W28:W39)", total_formula, report_total),
        ("AM22", "=SUM(AM10:AM21)", total_formula, report_total),
    )
    for cell, old, new, precise in overrides:
        if formulas.get(cell) != old:
            raise ValueError(f"Unknown V1 final total formula at Summary!{cell}.")
        cache = int(Decimal(str(precise)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        result.append(FinalTotalFormula(cell, old, new, cache))
    return tuple(result)


@dataclass(frozen=True)
class TemplateWritePlan:
    template_sheet: str
    summary_writes: tuple[CellWrite, ...]
    support_sheets: tuple[tuple[str, pd.DataFrame], ...]
    diagram_sheet_name: str
    diagram_rows: tuple[DiagramRow, ...]
    chart_caches: tuple[tuple[str, tuple[Any, ...], tuple[Any, ...]], ...]
    preserved_formulas: tuple[tuple[str, str], ...]
    skipped_formula_writes: tuple[tuple[str, str], ...]
    peak_formula_rows: tuple[PeakFormulaRow, ...]
    formula_bindings: tuple[FormulaBinding, ...]
    final_total_formulas: tuple[FinalTotalFormula, ...] = ()

    @property
    def ordered_writes(self) -> tuple[CellWrite, ...]:
        """Legacy COM order, including every support and Diagram_Data cell."""
        result: list[CellWrite] = []

        def append(sheet: str, cell: str, value: Any, source: str, force: bool = False) -> None:
            result.append(CellWrite(sheet, cell, value, source, len(result) + 1, force))

        for sheet, frame in self.support_sheets:
            if not len(frame.columns):
                append(sheet, "A1", "", f"support.{sheet}.empty")
                continue
            for col, name in enumerate(frame.columns, 1):
                append(sheet, f"{get_column_letter(col)}1", str(name), f"support.{sheet}.header.{name}")
            for row_number, values in enumerate(frame.itertuples(index=False, name=None), 2):
                for col, value in enumerate(values, 1):
                    append(sheet, f"{get_column_letter(col)}{row_number}", value, f"support.{sheet}.{row_number}.{col}")
        for col, value in enumerate(("movement_code", "total_pcu", "pm_peak_pcu", "am_peak_pcu"), 1):
            append(self.diagram_sheet_name, f"{get_column_letter(col)}1", value, "diagram.header")
        for item in self.diagram_rows:
            for col, value in enumerate((item.code, item.total_formula, item.pm_formula, item.am_formula), 1):
                append(self.diagram_sheet_name, f"{get_column_letter(col)}{item.row}", value, f"diagram.{item.code}.{col}")
        for write in self.summary_writes:
            append(write.sheet, write.cell, write.value, write.source, write.force)
        for binding in self.formula_bindings:
            append(self.template_sheet, binding.cell, binding.new_formula, f"formula_binding.{binding.period}.{binding.code}", True)
        return tuple(result)


def _in_ranges(cell: str, ranges: list[str]) -> bool:
    column, row, _, _ = range_boundaries(cell)
    for area in ranges:
        first_col, first_row, last_col, last_row = range_boundaries(str(area).split("!")[-1].replace("$", ""))
        if first_col <= column <= last_col and first_row <= row <= last_row:
            return True
    return False


def _range_cells(area: str) -> list[str]:
    first_col, first_row, last_col, last_row = range_boundaries(area.replace("$", ""))
    return [
        f"{get_column_letter(column)}{row}"
        for row in range(first_row, last_row + 1)
        for column in range(first_col, last_col + 1)
    ]


def diagram_data_formula(kind: str, code_ref: str, hourly_rows: int, hourly_columns: int) -> str:
    last = get_column_letter(max(hourly_columns, 2))
    data = f"'Hourly_Movement_PCU'!$B$2:${last}${hourly_rows}"
    times = f"'Hourly_Movement_PCU'!$A$2:$A${hourly_rows}"
    headers = f"'Hourly_Movement_PCU'!$B$1:${last}$1"
    match = f"MATCH({code_ref},{headers},0)"
    if kind == "total":
        return f"=IFERROR(INDEX({data},ROWS({times}),{match}),0)"
    peak = "'Peak_PHF'!$H$2" if kind == "pm" else "'Peak_PHF'!$D$2"
    key = f'IF(ISNUMBER({peak}),TEXT({peak},"hh:mm"),LEFT({peak},5))&"*"'
    return f"=IFERROR(INDEX({data},MATCH({key},{times},0),{match}),0)"


def resolve_template_write_plan(
    template_path: str | Path,
    template_map: dict[str, Any],
    report_data: dict[str, Any],
    metadata: dict[str, Any],
    chart_source_data: dict[str, Any],
    *,
    apply_v1_precision: bool = False,
) -> TemplateWritePlan:
    """Resolve the COM destination order into transport-independent cell writes."""

    template_sheet = str(template_map.get("template_sheet") or "Summary")
    workbook = load_workbook(template_path, read_only=True, data_only=False)
    try:
        summary = workbook[template_sheet]
        formulas = {
            cell.coordinate: str(cell.value)
            for row in summary.iter_rows()
            for cell in row
            if cell.data_type == "f"
        }
        original_times = {
            int(table["first_data_row"]): {
                row: summary[f"{table['time_column']}{row}"].value
                for row in range(int(table["first_data_row"]), int(table["total_row"]) + 1)
            }
            for table in (template_map["hourly_movement_table"], template_map["hourly_vehicle_class_table"])
        }
        helper_column = str(template_map["movement_diagram_cells"]["movement_value_rows"].get("helper_column") or "U")
        if helper_column != "U" or [summary[f"{helper_column}{row}"].value for row in range(9, 23)] != list(range(1, 15)):
            raise ValueError("Authoritative Summary!U9:U22 helper sequence must be 1..14.")
    finally:
        workbook.close()

    writable = list(template_map.get("writable_ranges") or [])
    overwritten = list(template_map.get("formula_overwrite_ranges") or [])
    protected = list(template_map.get("protected_formula_ranges") or [])
    writes: list[CellWrite] = []
    skipped: list[tuple[str, str]] = []

    def add(cell: str | None, value: Any, source: str, *, force: bool = False) -> None:
        if not cell:
            return
        if not force and not (_in_ranges(cell, writable) or _in_ranges(cell, overwritten)):
            if cell in formulas and _in_ranges(cell, protected):
                skipped.append((cell, source))
                return
            raise ValueError(f"Unmapped native-template write: {template_sheet}!{cell} ({source}).")
        if cell in formulas and not (force or _in_ranges(cell, overwritten)):
            skipped.append((cell, source))
            return
        writes.append(CellWrite(template_sheet, cell, value, source, len(writes) + 1, force))

    values = _metadata_values(metadata)
    for key, info in template_map.get("metadata_cells", {}).items():
        add(info.get("value_cell") or info.get("cell"), values.get(key, ""), f"metadata.{key}")
    movement = template_map.get("movement_diagram_cells", {})
    add(movement.get("diagram_title", {}).get("cell"), values.get("survey_point", ""), "movement.diagram_title")
    add(movement.get("diagram_date", {}).get("cell"), values.get("survey_date", ""), "movement.diagram_date")
    for key, info in movement.get("direction_labels", {}).items():
        add(info.get("cell"), _mapped_label_value(metadata, key), f"movement.direction_labels.{key}")
    for key, info in movement.get("road_labels", {}).items():
        add(info.get("cell"), _mapped_label_value(metadata, key), f"movement.road_labels.{key}")
    add(movement.get("caption", {}).get("cell"), metadata.get("caption_text") or "", "movement.caption")

    def peak_label(period: str) -> str:
        start = metadata.get(f"{period}_peak_start")
        end = metadata.get(f"{period}_peak_end")
        return f"{start}-{end}" if start and end else "-"

    if str(template_map.get("movement_code_scheme") or "") != "approach_movement":
        provenance = str(metadata.get("peak_selection_source") or "").strip() or "-"
        add("A3", f"Peak used for export: AM {peak_label('am')}; PM {peak_label('pm')} ({provenance})", "effective_peak_label", force=True)

    def add_table(table: dict[str, Any], frame: pd.DataFrame, source_name: str) -> None:
        first = int(table["first_data_row"])
        total = int(table["total_row"])
        time_column = str(table["time_column"])
        labels = original_times[first]
        source_time = _column_for_key(frame, "time")
        by_label: dict[str, pd.Series] = {}
        source_total: pd.Series | None = None
        if source_time:
            for _, source_row in frame.iterrows():
                label = _normalise_time_label(source_row[source_time])
                if label in {"รวม", "total"}:
                    source_total = source_row
                elif label:
                    by_label[label] = source_row
        for key, column in table.get("columns", {}).items():
            source_column = _column_for_key(frame, key)
            add(f"{column}{table['header_row']}", source_column or key, f"{source_name}.header.{key}")
            for row in range(first, total + 1):
                add(f"{column}{row}", "", f"{source_name}.clear.{key}")
            for row in range(first, total + 1):
                source_row = source_total if row == total else by_label.get(_normalise_time_label(labels[row]))
                if key == "time":
                    value = labels[row]
                elif source_row is not None and source_column is not None:
                    value = source_row[source_column]
                else:
                    value = 0
                add(f"{column}{row}", value, f"{source_name}.{key}.{row}")

    hourly = report_data.get("hourly_movement_pcu", pd.DataFrame())
    add_table(template_map["hourly_movement_table"], hourly, "hourly_movement_pcu")
    hourly_map = template_map["hourly_movement_table"]
    first, last = int(hourly_map["first_data_row"]), int(hourly_map["last_data_row"])
    time_labels = original_times[first]
    peak_rows: list[PeakFormulaRow] = []
    for match in _match_peak_rows(time_labels, hourly_map, helper_column, metadata):
        if not match.interval:
            raise ValueError(f"Effective {match.period.upper()} Peak interval is required for native formula binding.")
        if not match.matched or match.worksheet_row is None or match.helper_cell is None or match.helper_value is None:
            raise ValueError(f"Effective {match.period.upper()} Peak {match.interval!r} must match exactly one Summary time row.")
        start, end = match.interval.split("-", 1)
        peak_rows.append(PeakFormulaRow(match.period, start, end, match.worksheet_row, match.helper_cell, match.helper_value))

    movement_map = template_map["movement_diagram_cells"]
    movement_columns = [column for key, column in hourly_map["columns"].items() if key not in {"time", "Total"}]
    lookup_range = f"${movement_columns[0]}${hourly_map['header_row']}:${movement_columns[-1]}${hourly_map['total_row']}"
    bindings: list[FormulaBinding] = []
    mapped_codes: set[str] = set()
    for approach in movement_map["approach_tables"].values():
        for code in approach["movement_codes"]:
            if code in mapped_codes:
                raise ValueError(f"Duplicate native movement formula mapping for {code}.")
            mapped_codes.add(code)
            header_cell = approach["header_cells"][code]
            for peak in peak_rows:
                cell = approach[f"{peak.period}_peak_hour_cells"][code]
                old = formulas.get(cell, "")
                pattern = rf"=HLOOKUP\({re.escape(header_cell)},{re.escape(lookup_range)},\${helper_column}\$(\d+),FALSE\)"
                match = re.fullmatch(pattern, old)
                if match is None or not first <= int(match.group(1)) <= last:
                    raise ValueError(f"Unknown native HLOOKUP formula at Summary!{cell}: {old!r}.")
                new = old[:match.start(1)] + str(peak.worksheet_row) + old[match.end(1):]
                bindings.append(FormulaBinding(peak.period, code, cell, old, new))
            total_cell = approach["total_12_hour_cells"][code]
            total = f"=HLOOKUP({header_cell},{lookup_range},${helper_column}${hourly_map['total_row']},FALSE)"
            if formulas.get(total_cell) != total:
                raise ValueError(f"Unknown native total HLOOKUP formula at Summary!{total_cell}.")
    expected_codes = set(report_data.get("diagram_movement_codes") or template_map.get("movement_code_order") or MOVEMENT_CODES)
    if mapped_codes != expected_codes or len(bindings) != 32 or len({item.cell for item in bindings}) != 32:
        raise ValueError("Native Peak formula map must contain 16 unique movements and 32 AM/PM cells.")
    support = report_data.get("sheets") or {}
    vehicle = report_data.get("hourly_vehicle_class")
    if vehicle is None:
        vehicle = support.get("Hourly_Vehicle_Class", pd.DataFrame())
    add_table(template_map["hourly_vehicle_class_table"], vehicle, "hourly_vehicle_class")

    chart_sources = template_map.get("chart_anchors", {})
    for source_key, data_key in (("native_hourly_chart_source", "hourly_pcu"), ("native_vehicle_composition_chart_source", "vehicle_composition")):
        source = chart_sources.get(source_key, {})
        data = chart_source_data.get(data_key, {})
        for kind in ("categories", "values"):
            area = source.get(kind)
            if not area:
                continue
            refs = _range_cells(area)
            series = list(data.get(kind, []))
            if len(series) > len(refs):
                raise ValueError(f"Chart data exceeds mapped range {area}.")
            for index, ref in enumerate(refs):
                add(ref, series[index] if index < len(series) else "", f"chart.{data_key}.{kind}.{index}")

    codes = tuple(report_data.get("diagram_movement_codes") or template_map.get("movement_code_order") or MOVEMENT_CODES)
    diagram_values = movement_values_by_code(hourly, list(codes), metadata)
    diagram_rows = tuple(
        DiagramRow(
            row,
            str(code),
            diagram_data_formula("total", f"$A{row}", max(len(hourly) + 1, 2), max(len(hourly.columns), 2)),
            diagram_data_formula("pm", f"$A{row}", max(len(hourly) + 1, 2), max(len(hourly.columns), 2)),
            diagram_data_formula("am", f"$A{row}", max(len(hourly) + 1, 2), max(len(hourly.columns), 2)),
            float(diagram_values[str(code)]["total_12_hour"]),
            float(diagram_values[str(code)]["pm_peak"]),
            float(diagram_values[str(code)]["am_peak"]),
        )
        for row, code in enumerate(codes, 2)
    )
    if template_map.get("movement_code_scheme") == "approach_movement":
        final_totals = v2_final_total_formulas(formulas, support)
    elif apply_v1_precision:
        final_totals = v1_final_total_formulas(formulas, support)
    else:
        final_totals = ()
    rebound_cells = {binding.cell for binding in bindings} | {item.cell for item in final_totals}
    preserved = tuple(sorted((cell, formula) for cell, formula in formulas.items() if not _in_ranges(cell, overwritten) and cell not in rebound_cells))
    caches = tuple(
        (key, tuple(chart_source_data.get(key, {}).get("categories", ())), tuple(chart_source_data.get(key, {}).get("values", ())))
        for key in ("hourly_pcu", "vehicle_composition")
    )
    return TemplateWritePlan(
        template_sheet,
        tuple(writes),
        tuple((name, frame.copy(deep=True)) for name, frame in support.items()),
        str(report_data.get("diagram_data_sheet_name") or "Diagram_Data"),
        diagram_rows,
        caches,
        preserved,
        tuple(skipped),
        tuple(peak_rows),
        tuple(bindings),
        final_totals,
    )
