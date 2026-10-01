"""Standard routing reaches the production OOXML transport with reviewed Peaks."""

from __future__ import annotations

from datetime import time
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

import app
from tmc_processor.application.services import analyze_batch, analyze_single, assess_standard_peak_binding, export_batch_reviewed
from tmc_processor.batch import BatchItem, exclude_batch_item
from tmc_processor.diagram import MOVEMENT_CODES
from tmc_processor.exporter import EXCEL_TEMPLATE_EXPORT_MODE, SAFE_PNG_EXPORT_MODE, STANDARD_REPORT_EXPORT_MODE, PEAK_BINDING_FALLBACK_REASON
from tmc_processor.mapping_preset import load_mapping_preset
from tmc_processor import exporter, ooxml_template_export


def _inputs():
    raw_sheets = {}
    mapping_rows = []
    for index, code in enumerate(MOVEMENT_CODES):
        sheet_name = f"Movement {code}"
        rows = []
        for hour in range(7, 19):
            for minute in (0, 15, 30, 45):
                rows.append({
                    "raw_sheet": sheet_name,
                    "raw_direction": code,
                    "time_start": time(hour, minute),
                    "time_end": time(hour + 1, 0) if minute == 45 else time(hour, minute + 15),
                    "vehicle_class": "PC<7",
                    "count": index + hour - 5,
                })
        raw_sheets[sheet_name] = pd.DataFrame(rows)
        mapping_rows.append({
            "raw_sheet": sheet_name,
            "raw_direction": code,
            "movement_code": code,
            "source_stream": "mainline",
            "raw_movement_label": code,
            "from_leg": code[0],
            "to_leg": code[0] if code[1] == "U" else code[1],
            "turn_type": "U" if code[1] == "U" else "L" if code in {"NE", "SW", "WN", "ES"} else "T" if code in {"NS", "SN", "WE", "EW"} else "R",
            "facility_type": "at_grade",
            "include_in_peak": True,
            "include_in_report": True,
            "aggregation_method": "sum",
        })
    setup = {
        "project_name": "UX-7 production routing qualification",
        "survey_point": "All 16 movements",
        "survey_date_text": "2026-10-01",
        "survey_period": "07:00-19:00",
        "peak_selection_source": "user_confirmed",
    }
    return raw_sheets, pd.DataFrame(mapping_rows), setup


def _run(decision, *, ooxml: bool | None = None, am=("07:00", "08:00"), pm=("16:00", "17:00"), peak_mode="fixed_hourly"):
    raw_sheets, mapping, setup = _inputs()
    return analyze_single(
        raw_sheets=raw_sheets,
        mapping=mapping,
        setup=setup,
        detected_sheets=list(raw_sheets),
        confirmed_peak_periods={"AM": am, "PM": pm},
        peak_mode=peak_mode,
        generate_workbook=True,
        use_template_report_layout=decision.use_template_report_layout,
        use_excel_com_native_charts=decision.use_excel_com_native_charts,
        use_ooxml_native_template=decision.use_ooxml_native_template if ooxml is None else ooxml,
        export_mode=decision.backend_mode,
        export_mode_requested=decision.requested_mode,
        export_mode_used=decision.backend_mode,
        export_fallback_notice=decision.fallback_notice,
    )


def _fields(workbook):
    return dict(workbook["Export_Metadata"].iter_rows(min_row=2, values_only=True))


@pytest.mark.parametrize("com_available", [False, True])
def test_standard_application_path_uses_ooxml_with_or_without_com(monkeypatch, com_available):
    decision = app._standard_report_decision(SimpleNamespace(available=com_available, reason="probe result"))
    assert decision.backend_mode == EXCEL_TEMPLATE_EXPORT_MODE
    assert decision.use_ooxml_native_template and not decision.use_excel_com_native_charts
    original = ooxml_template_export.export_template_ooxml
    calls = []

    def record(*args, **kwargs):
        calls.append(args)
        return original(*args, **kwargs)

    monkeypatch.setattr(ooxml_template_export, "export_template_ooxml", record)
    monkeypatch.setattr(exporter, "_export_workbook_with_excel_com", lambda *a, **k: pytest.fail("COM used"))
    monkeypatch.setattr(exporter, "_export_workbook_from_template", lambda *a, **k: pytest.fail("PNG template used"))
    result = _run(decision, am=("08:00", "09:00"), pm=("17:00", "18:00"))
    assert len(calls) == 1
    with ZipFile(BytesIO(result.workbook_bytes)) as archive:
        assert len([name for name in archive.namelist() if name.startswith("xl/charts/chart") and name.endswith(".xml")]) == 2
        shape_count = sum(
            len(ET.fromstring(archive.read(name)).findall(".//{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}sp"))
            for name in archive.namelist() if name.startswith("xl/drawings/drawing") and name.endswith(".xml")
        )
        assert shape_count == 33
    workbook = load_workbook(BytesIO(result.workbook_bytes), read_only=True, data_only=False)
    cached = load_workbook(BytesIO(result.workbook_bytes), read_only=True, data_only=True)
    try:
        fields = _fields(cached)
        assert fields["export_mode_requested"] == STANDARD_REPORT_EXPORT_MODE
        assert fields["export_mode_used"] == EXCEL_TEMPLATE_EXPORT_MODE
        assert fields["effective_am_peak"] == "08:00-09:00"
        assert fields["effective_pm_peak"] == "17:00-18:00"
        assert [workbook["Summary"][f"U{row}"].value for row in range(9, 23)] == list(range(1, 15))
        assert "AM 08:00-09:00; PM 17:00-18:00" in workbook["Summary"]["A3"].value
        assert cached["Peak_PHF"]["D2"].value == "08:00"
        assert cached["Peak_PHF"]["H2"].value == "17:00"
        assert len([name for name in workbook.sheetnames if name == "Diagram_Data"]) == 1
    finally:
        workbook.close()
        cached.close()


def test_explicit_safe_png_application_path_never_calls_native_transport(monkeypatch):
    decision = SimpleNamespace(
        use_template_report_layout=False, use_excel_com_native_charts=False,
        use_ooxml_native_template=False, backend_mode=SAFE_PNG_EXPORT_MODE,
        requested_mode=SAFE_PNG_EXPORT_MODE, fallback_notice="",
    )
    monkeypatch.setattr(ooxml_template_export, "export_template_ooxml", lambda *a, **k: pytest.fail("OOXML used"))
    monkeypatch.setattr(exporter, "_export_workbook_with_excel_com", lambda *a, **k: pytest.fail("COM used"))
    result = _run(decision)
    workbook = load_workbook(BytesIO(result.workbook_bytes), read_only=True, data_only=True)
    try:
        assert _fields(workbook)["export_mode_used"] == SAFE_PNG_EXPORT_MODE
        assert "Summary" not in workbook.sheetnames
    finally:
        workbook.close()


def test_missing_template_decision_keeps_standard_request_and_safe_png(monkeypatch):
    decision = app._standard_report_decision(SimpleNamespace(available=False, reason="COM absent"), template_compatible=False)
    monkeypatch.setattr(ooxml_template_export, "export_template_ooxml", lambda *a, **k: pytest.fail("OOXML used"))
    result = _run(decision)
    workbook = load_workbook(BytesIO(result.workbook_bytes), read_only=True, data_only=True)
    try:
        fields = _fields(workbook)
        assert fields["export_mode_requested"] == STANDARD_REPORT_EXPORT_MODE
        assert fields["export_mode_used"] == SAFE_PNG_EXPORT_MODE
        assert "unavailable or incompatible" in fields["export_fallback_notice"]
    finally:
        workbook.close()


def test_ooxml_runtime_failure_has_explicit_safe_png_provenance(monkeypatch):
    decision = app._standard_report_decision(SimpleNamespace(available=False, reason="COM absent"))
    def fail(*args, **kwargs):
        raise OSError("forced transport failure")
    monkeypatch.setattr(ooxml_template_export, "export_template_ooxml", fail)
    monkeypatch.setattr(exporter, "_export_workbook_with_excel_com", lambda *a, **k: pytest.fail("COM used"))
    with pytest.warns(RuntimeWarning, match="OOXML native-template export failed"):
        result = _run(decision)
    workbook = load_workbook(BytesIO(result.workbook_bytes), read_only=True, data_only=True)
    try:
        fields = _fields(workbook)
        assert fields["export_mode_requested"] == STANDARD_REPORT_EXPORT_MODE
        assert fields["export_mode_used"] == SAFE_PNG_EXPORT_MODE
        assert "forced transport failure" in fields["export_fallback_notice"]
        assert "Summary" not in workbook.sheetnames
    finally:
        workbook.close()


def test_exact_peak_binding_accepts_hour_aligned_rolling_mode():
    binding = assess_standard_peak_binding({
        "am_peak_start": "08:00", "am_peak_end": "09:00",
        "pm_peak_start": "17:00", "pm_peak_end": "18:00",
        "peak_mode": "rolling_60min",
    })
    assert [(match.period, match.matched, match.worksheet_row, match.helper_value) for match in binding] == [
        ("am", True, 11, 3), ("pm", True, 20, 12),
    ]


def test_standard_non_aligned_confirmed_peak_uses_safe_png_without_ooxml(monkeypatch):
    binding = assess_standard_peak_binding({
        "am_peak_start": "08:15", "am_peak_end": "09:15",
        "pm_peak_start": "17:00", "pm_peak_end": "18:00",
    })
    assert not binding[0].matched and binding[1].matched
    decision = app._standard_report_decision(
        SimpleNamespace(available=True, reason="COM available"),
        template_compatible=False,
        fallback_reason=PEAK_BINDING_FALLBACK_REASON,
    )
    monkeypatch.setattr(ooxml_template_export, "export_template_ooxml", lambda *a, **k: pytest.fail("OOXML used"))
    monkeypatch.setattr(exporter, "_export_workbook_with_excel_com", lambda *a, **k: pytest.fail("COM used"))
    result = _run(decision, am=("08:15", "09:15"), pm=("17:00", "18:00"), peak_mode="rolling_60min")
    workbook = load_workbook(BytesIO(result.workbook_bytes), read_only=True, data_only=True)
    try:
        fields = _fields(workbook)
        assert fields["export_mode_requested"] == STANDARD_REPORT_EXPORT_MODE
        assert fields["export_mode_used"] == SAFE_PNG_EXPORT_MODE
        assert PEAK_BINDING_FALLBACK_REASON in fields["export_fallback_notice"]
        assert fields["effective_am_peak"] == "08:15-09:15"
        assert fields["effective_pm_peak"] == "17:00-18:00"
    finally:
        workbook.close()


def test_reviewed_batch_standard_uses_the_same_ooxml_transport(monkeypatch):
    root = Path(__file__).resolve().parents[1]
    demo = root / "samples" / "demo"
    source = demo / "DEMO_TMC1_FourLeg.xlsx"
    preset = load_mapping_preset((demo / "DEMO_TMC1_FourLeg.mapping.json").read_bytes()).preset
    analysis = analyze_batch(
        [BatchItem(file_name=source.name, workbook_bytes=source.read_bytes())],
        mapping_preset=preset,
        setup={"project_name": "Reviewed Batch routing"},
        generated_at="2026-10-01T00:00:00Z",
    )
    item = analysis.successful_items[0]
    item.confirmed_AM_peak = "08:00-09:00"
    item.confirmed_PM_peak = "17:00-18:00"
    decision = app._standard_report_decision(SimpleNamespace(available=False, reason="COM absent"))
    calls = []
    original = ooxml_template_export.export_template_ooxml
    def record(*args, **kwargs):
        calls.append(args)
        return original(*args, **kwargs)
    monkeypatch.setattr(ooxml_template_export, "export_template_ooxml", record)
    monkeypatch.setattr(exporter, "_export_workbook_with_excel_com", lambda *a, **k: pytest.fail("COM used"))
    result = export_batch_reviewed(
        analysis,
        setup={"project_name": "Reviewed Batch routing"},
        export_mode=decision.backend_mode,
        use_template_report_layout=True,
        use_excel_com_native_charts=False,
        use_ooxml_native_template=True,
        export_mode_requested=decision.requested_mode,
    )
    assert len(calls) == 1
    assert result.summary_rows[0].export_mode_requested == STANDARD_REPORT_EXPORT_MODE
    assert result.summary_rows[0].export_mode_used == EXCEL_TEMPLATE_EXPORT_MODE, result.summary_rows[0].notes
    with ZipFile(BytesIO(result.package_bytes)) as package:
        row = result.summary_rows[0]
        report = package.read(f"{row.folder_name}/{row.output_stem}_report.xlsx")
    with ZipFile(BytesIO(report)) as workbook:
        assert "xl/charts/chart1.xml" in workbook.namelist()
        assert "xl/charts/chart2.xml" in workbook.namelist()


def test_reviewed_batch_non_aligned_peak_falls_back_per_item(monkeypatch):
    root = Path(__file__).resolve().parents[1]
    demo = root / "samples" / "demo"
    source = demo / "DEMO_TMC1_FourLeg.xlsx"
    preset = load_mapping_preset((demo / "DEMO_TMC1_FourLeg.mapping.json").read_bytes()).preset
    analysis = analyze_batch(
        [BatchItem(file_name=source.name, workbook_bytes=source.read_bytes())],
        mapping_preset=preset, setup={"project_name": "Reviewed Batch routing"},
        generated_at="2026-10-01T00:00:00Z",
    )
    item = analysis.successful_items[0]
    item.confirmed_AM_peak = item.suggested_AM_peak
    item.confirmed_PM_peak = item.suggested_PM_peak
    assert item.confirmed_AM_peak == "08:15-09:15"
    monkeypatch.setattr(ooxml_template_export, "export_template_ooxml", lambda *a, **k: pytest.fail("OOXML used"))
    monkeypatch.setattr(exporter, "_export_workbook_with_excel_com", lambda *a, **k: pytest.fail("COM used"))
    result = export_batch_reviewed(
        analysis, setup={"project_name": "Reviewed Batch routing"},
        export_mode=EXCEL_TEMPLATE_EXPORT_MODE, use_template_report_layout=True,
        use_excel_com_native_charts=False, use_ooxml_native_template=True,
        export_mode_requested=STANDARD_REPORT_EXPORT_MODE,
    )
    row = result.summary_rows[0]
    assert row.export_status == "success"
    assert row.export_mode_requested == STANDARD_REPORT_EXPORT_MODE
    assert row.export_mode_used == SAFE_PNG_EXPORT_MODE
    assert row.confirmed_AM_peak == "08:15-09:15"
    assert PEAK_BINDING_FALLBACK_REASON in row.notes
    with ZipFile(BytesIO(result.package_bytes)) as package:
        report = package.read(f"{row.folder_name}/{row.output_stem}_report.xlsx")
    workbook = load_workbook(BytesIO(report), read_only=True, data_only=True)
    try:
        fields = _fields(workbook)
        assert fields["export_mode_requested"] == STANDARD_REPORT_EXPORT_MODE
        assert fields["export_mode_used"] == SAFE_PNG_EXPORT_MODE
        assert fields["export_fallback_notice"] == PEAK_BINDING_FALLBACK_REASON
        assert fields["effective_am_peak"] == "08:15-09:15"
    finally:
        workbook.close()


def test_reviewed_batch_mixes_native_fallback_excluded_and_failed_items(monkeypatch):
    demo = Path(__file__).resolve().parents[1] / "samples" / "demo"
    source_bytes = (demo / "DEMO_TMC1_FourLeg.xlsx").read_bytes()
    preset = load_mapping_preset((demo / "DEMO_TMC1_FourLeg.mapping.json").read_bytes()).preset
    analysis = analyze_batch(
        [
            BatchItem(file_name="aligned.xlsx", workbook_bytes=source_bytes),
            BatchItem(file_name="rolling.xlsx", workbook_bytes=source_bytes),
            BatchItem(file_name="excluded.xlsx", workbook_bytes=source_bytes),
            BatchItem(file_name="failed.xlsx", workbook_bytes=b"not an xlsx"),
        ],
        mapping_preset=preset, setup={"project_name": "Mixed reviewed Batch"},
        generated_at="2026-10-01T00:00:00Z",
    )
    by_name = {item.file_name: item for item in analysis.items}
    for name in ("aligned.xlsx", "rolling.xlsx"):
        by_name[name].confirmed_AM_peak = "08:00-09:00" if name == "aligned.xlsx" else "08:15-09:15"
        by_name[name].confirmed_PM_peak = "17:00-18:00"
    exclude_batch_item(by_name["excluded.xlsx"], "Operator excluded this file")
    calls = []
    original = ooxml_template_export.export_template_ooxml

    def record(*args, **kwargs):
        calls.append(args)
        return original(*args, **kwargs)

    monkeypatch.setattr(ooxml_template_export, "export_template_ooxml", record)
    monkeypatch.setattr(exporter, "_export_workbook_with_excel_com", lambda *a, **k: pytest.fail("COM used"))
    result = export_batch_reviewed(
        analysis, setup={"project_name": "Mixed reviewed Batch"},
        export_mode=EXCEL_TEMPLATE_EXPORT_MODE, use_template_report_layout=True,
        use_excel_com_native_charts=False, use_ooxml_native_template=True,
        export_mode_requested=STANDARD_REPORT_EXPORT_MODE,
    )
    rows = {row.file_name: row for row in result.summary_rows}
    assert len(calls) == 1
    assert (rows["aligned.xlsx"].export_status, rows["aligned.xlsx"].export_mode_used) == ("success", EXCEL_TEMPLATE_EXPORT_MODE)
    assert (rows["rolling.xlsx"].export_status, rows["rolling.xlsx"].export_mode_used) == ("success", SAFE_PNG_EXPORT_MODE)
    assert PEAK_BINDING_FALLBACK_REASON in rows["rolling.xlsx"].notes
    assert rows["excluded.xlsx"].export_status == "excluded"
    assert rows["failed.xlsx"].export_status == "failed"
    assert {row.export_mode_requested for row in rows.values()} == {STANDARD_REPORT_EXPORT_MODE}
    with ZipFile(BytesIO(result.package_bytes)) as package:
        names = set(package.namelist())
        assert f'{rows["aligned.xlsx"].folder_name}/{rows["aligned.xlsx"].output_stem}_report.xlsx' in names
        assert f'{rows["rolling.xlsx"].folder_name}/{rows["rolling.xlsx"].output_stem}_report.xlsx' in names
        assert not any(name.startswith(rows["excluded.xlsx"].folder_name + "/") for name in names)
        assert not any(name.startswith(rows["failed.xlsx"].folder_name + "/") for name in names)
