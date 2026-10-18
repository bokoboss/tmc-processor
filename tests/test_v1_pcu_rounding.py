"""V1 Native Summary final totals round precise sources once."""

from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
import pandas as pd
import pytest

from tmc_processor.importer import load_detected_sheets
from tmc_processor.mapping import read_mapping_excel
from tmc_processor.pipeline import process_tmc
from tmc_processor.exporter import _peak_report_frame
from tmc_processor.report_template import DEFAULT_TEMPLATE_PATH
from tmc_processor.report_template import DEFAULT_TEMPLATE_MAP_PATH, load_template_map
from tmc_processor.summaries import hourly_summary, movement_summary
from tmc_processor.template_write_plan import v1_final_total_formulas
from tmc_processor.template_write_verify import verify_ooxml_against_plan


ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "samples" / "demo"
AM_PM = {"AM": ("08:00", "09:00"), "PM": ("17:00", "18:00")}


def _round_once(value):
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _run(source="DEMO_TMC1_FourLeg.xlsx", *, contributions=None, pce=1.0,
         periods=AM_PM, report_excluded=(), peak_excluded=()):
    raw = load_detected_sheets(DEMO / source)
    mapping = read_mapping_excel(DEMO / "DEMO_TMC1_FourLeg_mapping.xlsx")
    if contributions is not None:
        for frame in raw.values():
            frame.loc[:, "count"] = 0
    for sheet, hour in contributions or ():
        candidates = raw[sheet].index[
            raw[sheet]["time_start"].astype(str).str.startswith(f"{hour:02}:00")
            & (raw[sheet]["vehicle_class"] == "PC<7")
        ]
        assert len(candidates)
        raw[sheet].loc[candidates[0], "count"] = 1
    if report_excluded:
        mapping.loc[mapping.raw_sheet.isin(report_excluded), "include_in_report"] = False
    if peak_excluded:
        mapping.loc[mapping.raw_sheet.isin(peak_excluded), "include_in_peak"] = False
    result = process_tmc(
        raw, mapping, {"project_name": "Issue 42 V1 precision"},
        detected_sheets=list(raw), confirmed_peak_periods=periods,
        pce_factors={"PC<7": pce}, use_template_report_layout=True,
        use_ooxml_native_template=True, generated_at="2026-10-01T00:00:00Z",
    )
    formulas = load_workbook(BytesIO(result.workbook_bytes), data_only=False)
    cached = load_workbook(BytesIO(result.workbook_bytes), data_only=True)
    return result, formulas, cached


def _assert_five_totals(result, formulas, cached):
    exact_report = math_fsum(result.normalized.loc[result.normalized.include_in_report, "pcu"])
    exact_am = float(result.peaks.iloc[0].hourly_pcu)
    exact_pm = float(result.peaks.iloc[1].hourly_pcu)
    expected = {"F30": _round_once(exact_am), "F31": _round_once(exact_pm),
                "F32": _round_once(exact_report), "W40": _round_once(exact_report),
                "AM22": _round_once(exact_report)}
    summary, values = formulas["Summary"], cached["Summary"]
    for ref, value in expected.items():
        assert values[ref].value == value, (ref, values[ref].value, value)
        assert summary[ref].data_type == "f"
    assert summary["F30"].value == "=ROUND('Peak_PHF'!$F$2,0)"
    assert summary["F31"].value == "=ROUND('Peak_PHF'!$J$2,0)"
    last = len(result.movement) + 1
    total_formula = f"=ROUND(SUMIFS('Movement_Summary'!$I$2:$I${last},'Movement_Summary'!$G$2:$G${last},TRUE),0)"
    for ref in ("F32", "W40", "AM22"):
        assert summary[ref].value == total_formula
    assert formulas["Summary"]["X40"].value == "=SUM(X28:X39)"
    assert formulas["Summary"]["AJ40"].value == "=SUM(AJ28:AJ39)"


def math_fsum(values):
    import math
    return math.fsum(float(value) for value in values)


@pytest.mark.parametrize("source,precise", [
    ("DEMO_TMC1_FourLeg.xlsx", 23959.041),
    ("DEMO_TMC1_FourLeg_Day2.xlsx", 23590.054),
])
def test_demo_final_cells_round_precise_survey_and_effective_peaks(source, precise):
    result, formulas, cached = _run(source)
    try:
        assert math_fsum(result.normalized.loc[result.normalized.include_in_report, "pcu"]) == pytest.approx(precise)
        _assert_five_totals(result, formulas, cached)
    finally:
        formulas.close()
        cached.close()


@pytest.mark.parametrize("contributions,pce,expected", [
    ([("ทิศ 1", 8), ("ทิศ 2", 8)], .49, 1),  # minimal 2 x .49
    ([("ทิศ 1", 8), ("ทิศ 2", 8), ("ทิศ 3", 8)], .49, 1),
    ([("ทิศ 1", 8), ("ทิศ 2", 8), ("ทิศ 3", 8)], .51, 2),
    ([("ทิศ 1", 8)], .5, 1),
    ([("ทิศ 1", 8), ("ทิศ 2", 8), ("ทิศ 3", 8)], .5, 2),
    ([("ทิศ 1", 7), ("ทิศ 1", 8), ("ทิศ 1", 9), ("ทิศ 2", 8), ("ทิศ 3", 8)], .5, 3),
])
def test_fractional_totals_use_excel_half_up_once(contributions, pce, expected):
    result, formulas, cached = _run(contributions=contributions, pce=pce)
    try:
        assert cached["Summary"]["F32"].value == expected
        _assert_five_totals(result, formulas, cached)
    finally:
        formulas.close()
        cached.close()


def test_multihour_accumulation_keeps_native_template_route():
    result, formulas, cached = _run(contributions=[("ทิศ 1", hour) for hour in (7, 8, 9)], pce=.49)
    try:
        _assert_five_totals(result, formulas, cached)
        assert cached["Summary"]["W10"].value == 0
        assert cached["Summary"]["W11"].value == 0
        assert cached["Summary"]["W12"].value == 0
    finally:
        formulas.close()
        cached.close()


def test_alternate_confirmed_am_pm_and_peak_exclusion_are_respected():
    result, formulas, cached = _run(
        contributions=[("ทิศ 1", 7), ("ทิศ 2", 16)], pce=.51,
        periods={"AM": ("07:00", "08:00"), "PM": ("16:00", "17:00")},
        peak_excluded=("ทิศ 1",),
    )
    try:
        _assert_five_totals(result, formulas, cached)
        assert cached["Summary"]["F30"].value == 0
        assert cached["Summary"]["F31"].value == 1
    finally:
        formulas.close()
        cached.close()


def test_report_excluded_source_is_absent_from_precise_final_total():
    result, formulas, cached = _run(
        contributions=[("ทิศ 1", 8), ("ทิศ 2", 8)], pce=.51,
        report_excluded=("ทิศ 1",),
    )
    try:
        assert math_fsum(result.normalized.loc[result.normalized.include_in_report, "pcu"]) == pytest.approx(.51)
        _assert_five_totals(result, formulas, cached)
    finally:
        formulas.close()
        cached.close()


def test_variable_movement_summary_rows_and_source_contract_rejections():
    raw = load_detected_sheets(DEMO / "DEMO_TMC1_FourLeg.xlsx")
    mapping = read_mapping_excel(DEMO / "DEMO_TMC1_FourLeg_mapping.xlsx")
    setup = {"am_peak_start": "08:00", "am_peak_end": "09:00",
             "pm_peak_start": "17:00", "pm_peak_end": "18:00",
             "peak_selection_source": "user_confirmed"}
    result = process_tmc(raw, mapping, setup, detected_sheets=list(raw),
                         confirmed_peak_periods=AM_PM, generate_workbook=False)
    normalized = result.normalized.copy()
    # Split one facility group in memory to exercise a different V1 row count.
    normalized.loc[normalized.index[0], "facility_type"] = "separated-test-group"
    movement = movement_summary(normalized)
    assert len(movement) == len(result.movement) + 1
    assert movement.groupby("movement_code").size().loc["NS"] >= 2
    support = {
        "Movement_Summary": movement,
        "Normalized_Data": normalized,
        "Hourly_Summary": hourly_summary(normalized),
        "Peak_PHF": _peak_report_frame(setup, result.peaks),
    }
    source = load_workbook(DEFAULT_TEMPLATE_PATH, read_only=True, data_only=False)
    try:
        summary = source["Summary"]
        formulas = {ref: summary[ref].value for ref in ("F30", "F31", "F32", "W40", "AM22")}
    finally:
        source.close()
    expected = v1_final_total_formulas(formulas, support)
    last = len(movement) + 1
    assert all(item.new_formula.endswith(f"$I${last},'Movement_Summary'!$G$2:$G${last},TRUE),0)")
               for item in expected if item.cell in {"F32", "W40", "AM22"})

    corrupted = {name: frame.copy() for name, frame in support.items()}
    corrupted["Movement_Summary"].loc[0, "pcu"] += 1
    with pytest.raises(ValueError, match="Movement_Summary"):
        v1_final_total_formulas(formulas, corrupted)
    corrupted = {name: frame.copy() for name, frame in support.items()}
    corrupted["Movement_Summary"] = corrupted["Movement_Summary"].rename(columns={"pcu": "pcu_source"})
    with pytest.raises(ValueError, match="schema"):
        v1_final_total_formulas(formulas, corrupted)
    altered = dict(formulas)
    altered["F32"] = "=SUM(AM22)"
    with pytest.raises(ValueError, match="F32"):
        v1_final_total_formulas(altered, support)
    corrupted = {name: frame.copy() for name, frame in support.items()}
    corrupted["Peak_PHF"].loc[0, "am_peak_pcu"] += 1
    with pytest.raises(ValueError, match="Peak_PHF"):
        v1_final_total_formulas(formulas, corrupted)


def _v1_contract_fixture():
    raw = load_detected_sheets(DEMO / "DEMO_TMC1_FourLeg.xlsx")
    mapping = read_mapping_excel(DEMO / "DEMO_TMC1_FourLeg_mapping.xlsx")
    setup = {"am_peak_start": "08:00", "am_peak_end": "09:00",
             "pm_peak_start": "17:00", "pm_peak_end": "18:00",
             "peak_selection_source": "user_confirmed"}
    result = process_tmc(raw, mapping, setup, detected_sheets=list(raw),
                         confirmed_peak_periods=AM_PM, generate_workbook=False)
    normalized = result.normalized.copy()
    movement = movement_summary(normalized)
    support = {
        "Movement_Summary": movement,
        "Normalized_Data": normalized,
        "Hourly_Summary": hourly_summary(normalized),
        "Peak_PHF": _peak_report_frame(setup, result.peaks),
    }
    source = load_workbook(DEFAULT_TEMPLATE_PATH, read_only=True, data_only=False)
    try:
        formulas = {ref: source["Summary"][ref].value for ref in ("F30", "F31", "F32", "W40", "AM22")}
    finally:
        source.close()
    return formulas, support


def test_rejects_balanced_per_hour_source_redistribution():
    formulas, support = _v1_contract_fixture()
    corrupted = {name: frame.copy() for name, frame in support.items()}
    hourly = corrupted["Hourly_Summary"]
    # These two demo hours remain on the same displayed integer after the
    # balanced shift, and the precise survey grand total remains unchanged.
    first, second = 0, 1
    before = hourly["pcu"].copy()
    displayed_before = [round(value) for value in before]
    hourly.loc[first, "pcu"] += 0.1
    hourly.loc[second, "pcu"] -= 0.1
    assert math_fsum(hourly["pcu"]) == pytest.approx(math_fsum(before), abs=1e-9)
    assert [round(value) for value in hourly["pcu"]] == displayed_before
    with pytest.raises(ValueError, match="hourly.*Normalized_Data|Normalized_Data.*hour"):
        v1_final_total_formulas(formulas, corrupted)


def test_rejects_duplicate_movement_summary_group_rows():
    formulas, support = _v1_contract_fixture()
    corrupted = {name: frame.copy() for name, frame in support.items()}
    movement = corrupted["Movement_Summary"]
    group_keys = ["movement_code", "from_leg", "to_leg", "turn_type", "facility_type",
                  "include_in_peak", "include_in_report"]
    index = movement["pcu"].astype(float).idxmax()
    duplicate = movement.loc[[index]].copy()
    movement.loc[index, "pcu"] = float(duplicate.iloc[0]["pcu"]) / 2
    duplicate.loc[:, "pcu"] = float(duplicate.iloc[0]["pcu"]) / 2
    corrupted["Movement_Summary"] = pd.concat([movement, duplicate], ignore_index=True)
    assert corrupted["Movement_Summary"].duplicated(subset=group_keys).any()
    assert math_fsum(corrupted["Movement_Summary"]["pcu"]) == pytest.approx(
        math_fsum(support["Movement_Summary"]["pcu"]), abs=1e-9
    )
    with pytest.raises(ValueError, match="duplicate.*group|group.*duplicate"):
        v1_final_total_formulas(formulas, corrupted)


def test_preserves_a_legitimate_zero_traffic_hour():
    formulas, support = _v1_contract_fixture()
    with_zero_hour = {name: frame.copy() for name, frame in support.items()}
    zero_hour = pd.DataFrame([{
        "hour_start": "06:00", "hour_end": "07:00", "count": 0, "pcu": 0.0,
    }], columns=["hour_start", "hour_end", "count", "pcu"])
    with_zero_hour["Hourly_Summary"] = pd.concat(
        [zero_hour, with_zero_hour["Hourly_Summary"]], ignore_index=True
    )
    result = v1_final_total_formulas(formulas, with_zero_hour)
    assert len(result) == 5


def test_native_verifier_rejects_corrupted_v1_final_cache(monkeypatch, tmp_path):
    import tmc_processor.template_write_plan as write_plan_module

    captured = []
    original = write_plan_module.resolve_template_write_plan

    def capture(*args, **kwargs):
        plan = original(*args, **kwargs)
        captured.append(plan)
        return plan

    monkeypatch.setattr(write_plan_module, "resolve_template_write_plan", capture)
    result, formulas, cached = _run(contributions=[("ทิศ 1", 8)], pce=.51)
    formulas.close()
    cached.close()
    assert len(captured) == 1
    broken = tmp_path / "corrupt-v1-cache.xlsx"
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(BytesIO(result.workbook_bytes)) as source, ZipFile(broken, "w") as output:
        for part in source.infolist():
            data = source.read(part.filename)
            if part.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                cells = {cell.attrib["r"]: cell for cell in root.iter(ns + "c")}
                cells["F32"].find(ns + "v").text = "0"
                data = ET.tostring(root)
            output.writestr(part, data)
    with pytest.raises(AssertionError, match="F32"):
        verify_ooxml_against_plan(DEFAULT_TEMPLATE_PATH, broken,
                                  load_template_map(DEFAULT_TEMPLATE_MAP_PATH), captured[0])
