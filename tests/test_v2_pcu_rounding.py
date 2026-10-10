"""Final V2 totals must not accumulate rounding from displayed hourly cells."""
from dataclasses import replace
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from test_phase_f_v2_dry_run import (
    _dry_run_with_preset, _setup, _raw_sheets, _v2_preset_mapping,
    V2_TEMPLATE_WORKBOOK,
)
from tmc_processor.exporter import export_v2_generated_workbook
from tmc_processor.peaks import confirmed_peak_phf
from tmc_processor.summaries import hourly_movement_pcu, hourly_summary, movement_summary
from tmc_processor.exporter import _v2_template_export_sheets, _native_chart_source_data
from tmc_processor.report_template import load_template_map
from tmc_processor.template_write_plan import resolve_template_write_plan, v2_final_total_formulas
from tmc_processor.template_write_verify import verify_ooxml_against_plan
from test_phase_f_v2_dry_run import V2_TEMPLATE_MAP


def _confirmed_setup():
    return {**_setup(), "am_peak_start": "08:00", "am_peak_end": "09:00",
            "pm_peak_start": "17:00", "pm_peak_end": "18:00",
            "peak_selection_source": "user_confirmed"}


def _round_once(value):
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _assert_totals(result, expected_total, expected_am, expected_pm):
    payload = export_v2_generated_workbook(result, setup=_confirmed_setup())
    cached = load_workbook(BytesIO(payload), data_only=True)
    formulas = load_workbook(BytesIO(payload), data_only=False)
    for ref in ("F32", "W40", "AM22"):
        assert cached["Summary"][ref].value == expected_total
        assert formulas["Summary"][ref].value == "=ROUND(SUM('Movement_Summary'!$G$2:$G$17),0)"
    for ref, source, expected in (("F30", "F", expected_am), ("F31", "J", expected_pm)):
        assert cached["Summary"][ref].value == expected
        assert formulas["Summary"][ref].value == f"=ROUND('Peak_PHF'!${source}$2,0)"
    for name, frame, column in (("Normalized_Data", result.normalized, "pcu"),
                                ("Movement_Summary", result.movement, "pcu"),
                                ("Hourly_Totals", result.hourly, "pcu")):
        rows = list(cached[name].values)
        index = rows[0].index(column)
        assert sum(row[index] for row in rows[1:]) == pytest.approx(frame[column].sum())
    assert formulas.calculation.fullCalcOnLoad and formulas.calculation.forceFullCalc
    with ZipFile(BytesIO(payload)) as output, ZipFile(V2_TEMPLATE_WORKBOOK) as template:
        for part in template.namelist():
            if part.startswith(("xl/drawings/", "xl/media/")):
                assert output.read(part) == template.read(part)
    cached.close()
    formulas.close()
    return payload


def test_demo_round_once_totals_and_confirmed_peaks():
    result = _dry_run_with_preset()
    assert result.normalized.pcu.sum() == pytest.approx(23959.041)
    _assert_totals(result, 23959, 2981, 3153)


@pytest.mark.parametrize("fraction,single", [(0.49, False), (0.51, False), (0.5, False),
                                            (2.5, False), (-0.5, False), (0.5, True),
                                            (2.5, True), (-0.5, True)])
def test_fractional_hourly_and_movement_rounding_does_not_accumulate(fraction, single):
    result = _dry_run_with_preset()
    normalized = result.normalized.copy()
    normalized["pcu"] = 0.0
    # One fractional contribution per movement per hour, including .5 ties.
    indices = normalized.groupby(["movement_code", normalized.time_start.astype(str).str[:2]]).head(1).index
    if single:
        indices = normalized[normalized.time_start.astype(str).str.startswith("08:")].head(1).index
    normalized.loc[indices, "pcu"] = fraction
    mapping = _v2_preset_mapping(_raw_sheets())
    peaks = confirmed_peak_phf(normalized, {"AM": ("08:00", "09:00"), "PM": ("17:00", "18:00")})
    result = replace(result, normalized=normalized, movement=movement_summary(normalized),
                     hourly=hourly_summary(normalized),
                     hourly_movement_pcu=hourly_movement_pcu(normalized, mapping, "approach_movement"),
                     peaks=peaks)
    exact = normalized.pcu.sum()
    if not single:
        assert sum(_round_once(value) for value in result.hourly.pcu) != _round_once(exact)
    _assert_totals(result, _round_once(exact), _round_once(peaks.iloc[0].hourly_pcu),
                   _round_once(peaks.iloc[1].hourly_pcu))


@pytest.fixture(scope="module")
def demo_plan():
    result = _dry_run_with_preset()
    setup = _confirmed_setup()
    sheets, hourly, vehicle, _ = _v2_template_export_sheets(result, setup, None)
    mapping = load_template_map(V2_TEMPLATE_MAP)
    data = {"sheets": sheets, "hourly_movement_pcu": hourly,
            "diagram_movement_codes": sheets["Movement_Summary"].movement_code.tolist(),
            "diagram_data_sheet_name": "Movement_Diagram_Data"}
    plan = resolve_template_write_plan(V2_TEMPLATE_WORKBOOK, mapping, data, setup,
                                       _native_chart_source_data(hourly, vehicle))
    return result, mapping, plan


@pytest.mark.parametrize("fault", ["original_formula", "movement_order", "normalized", "hourly", "nonfinite"])
def test_final_total_override_rejects_unvalidated_sources(demo_plan, fault):
    _, _, plan = demo_plan
    formulas = {item.cell: item.old_formula for item in plan.final_total_formulas}
    support = {name: frame.copy() for name, frame in plan.support_sheets}
    if fault == "original_formula":
        formulas["W40"] = "=SUM(W28:W38)"
    elif fault == "movement_order":
        support["Movement_Summary"] = support["Movement_Summary"].iloc[::-1]
    elif fault in {"normalized", "hourly"}:
        name = "Normalized_Data" if fault == "normalized" else "Hourly_Totals"
        support[name].loc[0, "pcu"] += 1
    else:
        support["Peak_PHF"].loc[0, "am_peak_pcu"] = float("nan")
    with pytest.raises(ValueError):
        v2_final_total_formulas(formulas, support)


@pytest.mark.parametrize("fault", ["formula", "cache", "vehicle_formula"])
def test_verifier_rejects_corrupted_v2_final_totals(demo_plan, tmp_path, fault):
    from xml.etree import ElementTree as ET
    result, mapping, plan = demo_plan
    payload = export_v2_generated_workbook(result, setup=_confirmed_setup())
    broken = tmp_path / "broken.xlsx"
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(BytesIO(payload)) as source, ZipFile(broken, "w") as output:
        for part in source.infolist():
            data = source.read(part.filename)
            if part.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                cells = {cell.attrib["r"]: cell for cell in root.iter(ns + "c")}
                if fault == "cache":
                    cells["W40"].find(ns + "v").text = "23956"
                elif fault == "formula":
                    cells["W40"].find(ns + "f").text = "SUM(W28:W39)"
                else:
                    cells["X40"].find(ns + "f").text = "SUM(X28:X38)"
                data = ET.tostring(root)
            output.writestr(part, data)
    with pytest.raises(AssertionError, match="W40|X40"):
        verify_ooxml_against_plan(V2_TEMPLATE_WORKBOOK, broken, mapping, plan)
