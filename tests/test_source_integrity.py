"""Source defects must remain visible after movement/peak aggregation."""

from datetime import time
from io import BytesIO
from zipfile import ZipFile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from tmc_processor.importer import parse_raw_count_sheet
from tmc_processor.pipeline import process_tmc


WINDOWS = {"AM": ("07:00", "08:00"), "PM": ("07:00", "08:00")}
CONFIRMED = {"AM": ("07:00", "08:00"), "PM": ("07:00", "08:00")}


def mapping():
    return pd.DataFrame([
        dict(raw_sheet=sheet, raw_direction=str(index + 1), movement_code="NS",
             from_leg="N", to_leg="S", turn_type="through", facility_type="at_grade",
             source_stream=stream, include_in_peak=True, include_in_report=True,
             aggregation_method="sum")
        for index, (sheet, stream) in enumerate([("ทิศ 1", "mainline"), ("ทิศ 2", "frontage")])
    ])


def vendor_table():
    return pd.DataFrame([
        ["start", "end", "รถยนต์นั่งส่วนบุคคลไม่เกิน 7 คน", "รถจักรยานยนต์"],
        *[[time(7, minute), time(8) if minute == 45 else time(7, minute + 15), 10, 2]
          for minute in (0, 15, 30, 45)],
    ])


def vendor_sources(defect=None):
    first = vendor_table()
    if defect == "missing":
        first = first.drop(index=2)
    elif defect == "duplicate":
        first = pd.concat([first, first.iloc[[2]]], ignore_index=True)
    elif defect == "duration":
        first.iat[2, 1] = time(7, 35)
    elif defect == "misaligned":
        first.iat[2, 0], first.iat[2, 1] = time(7, 16), time(7, 31)
    elif defect == "blank":
        first.iat[2, 2] = None
    elif defect == "invalid":
        first.iat[2, 2] = "bad count"
    elif defect == "zero":
        first.iat[2, 2] = 0
    return {"ทิศ 1": first, "ทิศ 2": vendor_table()}


def sources(defect=None):
    return {sheet: parse_raw_count_sheet(table, sheet).data
            for sheet, table in vendor_sources(defect).items()}


def process(raw, confirmed=False, export=False):
    return process_tmc(raw, mapping(), {}, peak_windows=WINDOWS,
                       confirmed_peak_periods=CONFIRMED if confirmed else None,
                       generate_workbook=export)


@pytest.mark.parametrize("confirmed", [False, True], ids=["auto", "confirmed"])
@pytest.mark.parametrize("defect,check,severity", [
    ("missing", "source_missing_interval", "error"),
    ("duplicate", "source_duplicate_interval", "error"),
    ("duration", "source_invalid_interval_duration", "error"),
    ("misaligned", "source_misaligned_interval", "error"),
    ("blank", "source_blank_count", "warning"),
    ("invalid", "source_invalid_count", "error"),
])
def test_source_defect_is_visible_for_both_peak_paths(defect, check, severity, confirmed):
    result = process(sources(defect), confirmed)
    issues = result.qc[result.qc.check == check]
    assert not issues.empty, result.qc.to_string()
    assert set(issues.raw_sheet) == {"ทิศ 1"}
    assert set(issues.severity) == {severity}
    assert issues.detail.str.contains("07:15" if defect != "misaligned" else "07:16").all()
    if defect != "missing":
        assert issues.detail.str.contains("row").all()
    assert not result.peaks.empty  # advisory QC must not change the peak calculation contract


def test_blank_zero_and_invalid_have_distinct_source_evidence():
    for defect, status in [("blank", "blank"), ("zero", "zero"), ("invalid", "invalid")]:
        raw = sources(defect)["ทิศ 1"]
        row = raw[(raw.time_start == time(7, 15)) & (raw.vehicle_class == "PC<7")].iloc[0]
        assert row["source_count_status"] == status
        assert row["source_row"] == 3
        assert row["source_column"] == 3
        if defect == "invalid":
            assert row["source_count_raw"] == "bad count"
    result = process(sources("zero"))
    assert not result.qc.check.isin(["source_blank_count", "source_invalid_count"]).any()


def test_valid_classes_and_intentional_stream_aggregation_are_not_duplicates():
    result = process(sources())
    assert not result.qc.check.str.startswith("source_").any()
    assert result.normalized["count"].sum() == 96
    assert result.normalized.pcu.sum() == pytest.approx(85.328)
    assert result.peaks.hourly_pcu.tolist() == pytest.approx([85.328, 85.328])
    assert result.peaks.phf.tolist() == [1.0, 1.0]


def test_single_qc_display_and_export_retain_source_provenance():
    from tmc_processor.ui.app_shell import _qc_display_frame

    result = process(sources("invalid"), export=True)
    display = _qc_display_frame(result.qc, thai_labels=False)
    issue = display[display.check == "source_invalid_count"].iloc[0]
    assert issue.raw_sheet == "ทิศ 1"
    assert "C3" in issue.detail and "bad count" in issue.detail
    workbook = load_workbook(BytesIO(result.workbook_bytes), read_only=True, data_only=True)
    rows = list(workbook["QC_Check"].values)
    records = [dict(zip(rows[0], row)) for row in rows[1:]]
    exported = next(row for row in records if row["check"] == "source_invalid_count")
    assert exported["detail"] == issue.detail
    assert exported["raw_sheet"] == issue.raw_sheet


@pytest.mark.parametrize("defect,check", [
    ("missing", "source_missing_interval"), ("duplicate", "source_duplicate_interval"),
    ("duration", "source_invalid_interval_duration"), ("misaligned", "source_misaligned_interval"),
    ("blank", "source_blank_count"), ("invalid", "source_invalid_count"),
])
def test_batch_analysis_confirmation_and_export_keep_source_checks(defect, check):
    from tmc_processor.batch import (BatchItem, analyze_batch_files, batch_analysis_qc_rows,
                                     generate_batch_zip_from_reviewed_peaks)
    from tmc_processor.ui.app_shell import _batch_qc_preview_display_frame

    payload = BytesIO()
    with pd.ExcelWriter(payload, engine="openpyxl") as writer:
        for sheet, table in vendor_sources(defect).items():
            table.to_excel(writer, sheet_name=sheet, header=False, index=False)
    analysis = analyze_batch_files([BatchItem("source.xlsx", payload.getvalue())],
                                   mapping=mapping(), setup={}, peak_windows=WINDOWS)
    item = analysis.items[0]
    assert item.status == "success", item.notes
    assert check in set(item.qc.check)
    display = _batch_qc_preview_display_frame(pd.DataFrame(batch_analysis_qc_rows(analysis)))
    assert check in set(display["รายการตรวจ"])
    item.confirmed_AM_peak = item.confirmed_PM_peak = "07:00-08:00"
    result = generate_batch_zip_from_reviewed_peaks(analysis, setup={}, peak_windows=WINDOWS,
                                                   use_template_report_layout=False)
    summary = result.summary_rows[0]
    assert summary.status == "success", summary.notes
    assert summary.QC_errors > 0 if defect != "blank" else summary.QC_warnings > 0
    assert check in {row["check"] for row in result.qc_rows}
    with ZipFile(BytesIO(result.package_bytes)) as archive:
        workbook = load_workbook(BytesIO(archive.read(summary.export_file)), read_only=True, data_only=True)
        rows = list(workbook["QC_Check"].values)
        records = [dict(zip(rows[0], row)) for row in rows[1:]]
        assert any(row["check"] == check and row["raw_sheet"] == "ทิศ 1" for row in records)
        summary_book = load_workbook(BytesIO(archive.read("batch_summary.xlsx")), read_only=True, data_only=True)
        qc_rows = list(summary_book["Batch_QC"].values)
        assert check in [dict(zip(qc_rows[0], row))["check"] for row in qc_rows[1:]]


@pytest.mark.parametrize("defect,check", [("missing", "source_missing_interval"),
                                         ("duplicate", "source_duplicate_interval"),
                                         ("duration", "source_invalid_interval_duration"),
                                         ("misaligned", "source_misaligned_interval"),
                                         ("blank", "source_blank_count"),
                                         ("invalid", "source_invalid_count")])
def test_wide_tables_preserve_the_same_integrity_checks(defect, check):
    wide = {}
    for sheet, table in vendor_sources(defect).items():
        table = table.iloc[1:].copy()
        table.columns = ["time_start", "time_end", "PC<7", "MC"]
        wide[sheet] = table
    result = process(wide)
    assert check in set(result.qc.check)


def test_multiple_columns_of_one_vehicle_class_are_one_interval_row():
    table = vendor_table()
    table[4] = ["รถจักรยานยนต์", 1, 1, 1, 1]
    raw = sources()
    raw["ทิศ 1"] = parse_raw_count_sheet(table, "ทิศ 1").data
    result = process(raw)
    assert "source_duplicate_interval" not in set(result.qc.check)


def test_legacy_long_tables_without_provenance_detect_duplicate_class_records():
    raw = {sheet: frame.iloc[:, :6].copy() for sheet, frame in sources("duplicate").items()}
    assert "source_duplicate_interval" in set(process(raw).qc.check)
    raw = {sheet: frame.iloc[:, :6].copy() for sheet, frame in sources().items()}
    assert "source_duplicate_interval" not in set(process(raw).qc.check)


def test_repeated_source_cell_records_also_inflate_counts_and_are_reported():
    raw = sources()
    raw["ทิศ 1"] = pd.concat([raw["ทิศ 1"], raw["ทิศ 1"].iloc[[2]]], ignore_index=True)
    assert "source_duplicate_interval" in set(process(raw).qc.check)


def test_qc_does_not_fan_out_source_issues_with_mapping_rows():
    from tmc_processor.qc import run_qc

    raw = sources("duplicate")
    result = process(raw)
    qc = run_qc(list(raw), pd.concat([mapping(), mapping()]), result.normalized,
                peaks=result.peaks, peak_windows=WINDOWS, raw_sheets=raw)
    assert len(qc[qc.check == "source_duplicate_interval"]) == 1


def test_excluded_report_source_does_not_impose_coverage_on_included_sources():
    raw = sources()
    raw["ทิศ 2"].loc[raw["ทิศ 2"].time_start == time(7), "time_start"] = time(6, 45)
    raw["ทิศ 2"].loc[raw["ทิศ 2"].time_end == time(7, 15), "time_end"] = time(7)
    selected = mapping()
    selected.loc[selected.raw_sheet == "ทิศ 2", "include_in_report"] = False
    result = process_tmc(raw, selected, {}, peak_windows=WINDOWS, generate_workbook=False)
    assert "source_missing_interval" not in set(result.qc.check)


def test_confirmed_peak_requires_each_included_source_even_if_all_omit_a_slot():
    raw = sources("missing")
    raw["ทิศ 2"] = raw["ทิศ 2"][raw["ทิศ 2"].time_start != time(7, 15)]
    result = process(raw, confirmed=True)
    missing = result.qc[result.qc.check == "source_missing_interval"]
    assert set(missing.raw_sheet) == set(raw)
    assert missing.detail.str.contains("affected_peaks=AM,PM").all()


@pytest.mark.parametrize("value", [" ", "0", "not-a-count", float("inf")])
def test_wide_count_states_are_not_lost_before_qc(value):
    raw = {sheet: pd.DataFrame(dict(time_start=[time(7, m) for m in (0, 15, 30, 45)],
                                   time_end=[time(7, m) for m in (15, 30, 45)] + [time(8)],
                                   **{"PC<7": [10, value, 10, 10]})) for sheet in sources()}
    result = process(raw)
    expected = "source_blank_count" if value == " " else (None if value == "0" else "source_invalid_count")
    if expected:
        assert expected in set(result.qc.check)
    else:
        assert not result.qc.check.isin(["source_blank_count", "source_invalid_count"]).any()


def test_native_template_export_retains_actionable_source_qc():
    from tmc_processor.importer import load_detected_sheets, load_detected_sheet_details
    from tmc_processor.mapping import read_mapping_excel

    demo = Path(__file__).resolve().parents[1] / "samples/demo"
    details = load_detected_sheet_details(demo / "DEMO_TMC1_FourLeg.xlsx")
    sheet, parsed = next(iter(details.items()))
    column = next(item["column"] for item in parsed.debug.vehicle_class_columns if item["vehicle_class"] == "PC<7")
    workbook = load_workbook(demo / "DEMO_TMC1_FourLeg.xlsx")
    workbook[sheet].cell(parsed.debug.first_data_row, column).value = "bad count"
    payload = BytesIO()
    workbook.save(payload)
    result = process_tmc(load_detected_sheets(BytesIO(payload.getvalue())),
                         read_mapping_excel(demo / "DEMO_TMC1_FourLeg_mapping.xlsx"), {},
                         confirmed_peak_periods={"AM": ("07:00", "08:00"), "PM": ("17:00", "18:00")},
                         use_template_report_layout=True, use_ooxml_native_template=True)
    output = load_workbook(BytesIO(result.workbook_bytes), read_only=True, data_only=True)
    rows = list(output["QC_Check"].values)
    records = [dict(zip(rows[0], row)) for row in rows[1:]]
    issue = next(row for row in records if row["check"] == "source_invalid_count")
    assert issue["raw_sheet"] == sheet and "bad count" in issue["detail"]
    with ZipFile(BytesIO(result.workbook_bytes)) as package:
        assert "xl/charts/chart1.xml" in package.namelist()


@pytest.mark.parametrize("token", ["NA", "N/A", "NULL", "nan"])
def test_excel_text_na_markers_are_invalid_counts_not_blank_cells(token):
    from tmc_processor.importer import load_detected_sheets

    table = vendor_table()
    table.iat[2, 2] = token
    payload = BytesIO()
    with pd.ExcelWriter(payload, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="ทิศ 1", header=False, index=False)
    raw = load_detected_sheets(BytesIO(payload.getvalue()))["ทิศ 1"]
    row = raw[(raw.time_start == time(7, 15)) & (raw.vehicle_class == "PC<7")].iloc[0]
    assert row.source_count_status == "invalid"
    assert row.source_count_raw == token
