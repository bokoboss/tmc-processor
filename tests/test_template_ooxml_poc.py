from __future__ import annotations

import hashlib
import warnings
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from scripts.inspect_template_package import inventory_package, patch_one_cell, worksheet_part_for_name
from tmc_processor.report_template import DEFAULT_TEMPLATE_MAP_PATH, DEFAULT_TEMPLATE_PATH, load_template_map
from tmc_processor.template_audit import iter_formula_cells


def _parts(path: Path) -> dict[str, bytes]:
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _load(path: Path):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="DrawingML support is incomplete.*", category=UserWarning)
        return load_workbook(path, data_only=False)


def test_one_cell_ooxml_patch_preserves_every_other_template_part(tmp_path: Path) -> None:
    source = DEFAULT_TEMPLATE_PATH
    output = tmp_path / "one_cell_ooxml_proof.xlsx"
    value = "OOXML POC - metadata value"
    mapping = load_template_map(DEFAULT_TEMPLATE_MAP_PATH)
    with ZipFile(source) as archive:
        sheet_part = worksheet_part_for_name(archive, mapping["template_sheet"])
        original_names = archive.namelist()
    before = _parts(source)
    source_sha = hashlib.sha256(source.read_bytes()).hexdigest()

    inventory = inventory_package(source, writable_parts=(sheet_part,))
    assert inventory == inventory_package(source, writable_parts=(sheet_part,))
    assert [part.path for part in inventory] == sorted(original_names)
    assert {part.path for part in inventory if part.category == "expected writable"} == {sheet_part}
    assert not any(part.category == "expected structural change" for part in inventory)
    assert all(part.size == len(before[part.path]) for part in inventory)
    assert all(part.sha256 == hashlib.sha256(before[part.path]).hexdigest() for part in inventory)

    assert patch_one_cell(source, output, sheet_name="Summary", cell_ref="E5", value=value) == sheet_part
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_sha
    with ZipFile(output) as archive:
        assert archive.testzip() is None
        assert archive.namelist() == original_names
    after = _parts(output)
    assert {name for name in before if before[name] != after[name]} == {sheet_part}
    assert all(after[name] == before[name] for name in before if name.endswith(".rels"))
    assert after["[Content_Types].xml"] == before["[Content_Types].xml"]
    assert after["xl/styles.xml"] == before["xl/styles.xml"]
    for prefix in ("xl/charts/", "xl/drawings/", "xl/media/", "xl/embeddings/", "xl/theme/"):
        assert all(after[name] == before[name] for name in before if name.startswith(prefix))

    assert iter_formula_cells(output) == iter_formula_cells(source)
    original, patched = _load(source), _load(output)
    try:
        assert patched.sheetnames == original.sheetnames
        old_sheet, new_sheet = original["Summary"], patched["Summary"]
        assert new_sheet["E5"].value == value
        assert new_sheet["E5"].style_id == old_sheet["E5"].style_id
        assert new_sheet.calculate_dimension() == old_sheet.calculate_dimension()
        assert {str(r) for r in new_sheet.merged_cells.ranges} == {str(r) for r in old_sheet.merged_cells.ranges}
        assert new_sheet.print_area == old_sheet.print_area
    finally:
        original.close()
        patched.close()

    result = inventory_package(output, writable_parts=(sheet_part,))
    assert {part.path for part in result if part.category == "must remain byte-identical"} == set(before) - {sheet_part}