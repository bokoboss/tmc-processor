from __future__ import annotations
import hashlib
import warnings
import zipfile
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET
import pandas as pd
import pytest
from openpyxl import load_workbook
from scripts.inspect_template_package import compare_package_parts, inventory_package, summarize_package_diff
from tmc_processor.ooxml_template_export import export_template_ooxml, _put_summary
from tmc_processor.exporter import _native_chart_source_data
from tmc_processor.report_template import DEFAULT_TEMPLATE_MAP_PATH, DEFAULT_TEMPLATE_PATH, load_template_map
from tmc_processor.template_audit import audit_template, iter_formula_cells
from tmc_processor.diagram import MOVEMENT_CODES
from tmc_processor.template_write_plan import resolve_template_write_plan
from tmc_processor.template_write_verify import verify_ooxml_against_plan
from tmc_processor import excel_com_export

MAIN="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
CHART="http://schemas.openxmlformats.org/drawingml/2006/chart"
PKG_REL="http://schemas.openxmlformats.org/package/2006/relationships"
DOC_REL="http://schemas.openxmlformats.org/officeDocument/2006/relationships"

def _payload():
    times=[f"{hour:02}:00-{hour+1:02}:00" for hour in range(7,19)]
    movement_codes=["NE","NS","NW","NU","SW","SN","SE","SU","WN","WE","WS","WU","ES","EW","EN","EU"]
    hourly_values=[[i+j for j in range(16)] for i in range(1,13)]
    movement_totals=[sum(row[j] for row in hourly_values) for j in range(16)]
    hourly=pd.DataFrame([[label,*values,sum(values)] for label,values in zip(times,hourly_values)]+[["Total",*movement_totals,sum(movement_totals)]],columns=["time",*movement_codes,"Total"])
    leg_pairs={"NE":("N","E","L"),"NS":("N","S","T"),"NW":("N","W","R"),"NU":("N","N","U"),"SW":("S","W","L"),"SN":("S","N","T"),"SE":("S","E","R"),"SU":("S","S","U"),"WN":("W","N","L"),"WE":("W","E","T"),"WS":("W","S","R"),"WU":("W","W","U"),"ES":("E","S","L"),"EW":("E","W","T"),"EN":("E","N","R"),"EU":("E","E","U")}
    mapping_frame=pd.DataFrame([{"raw_sheet":f"Sheet {i+1}","raw_direction":str(i+1),"movement_code":code,"from_leg":leg_pairs[code][0],"to_leg":leg_pairs[code][1],"turn_type":leg_pairs[code][2],"include_in_report":True} for i,code in enumerate(movement_codes)])
    normalized_frame=pd.DataFrame([{"from_leg":leg_pairs[code][0],"to_leg":leg_pairs[code][1],"turn_type":leg_pairs[code][2],"output_movement_code":code,"include_in_report":True,"include_in_peak":True,"time_start":label[:5],"pcu":hourly_values[i][j]} for i,label in enumerate(times) for j,code in enumerate(movement_codes)])
    movement_summary=pd.DataFrame([{"movement_code":code,"from_leg":leg_pairs[code][0],"to_leg":leg_pairs[code][1],"turn_type":leg_pairs[code][2],"include_in_report":True,"pcu":movement_totals[j]} for j,code in enumerate(movement_codes)])
    vehicle_classes=["Bicy","MC","PC<7","PC>7","LB","MB","HB","LT","MT","HT","TR","STR"]
    vehicle=pd.DataFrame([[label,100+i,50+i,*([i]*12)] for i,label in enumerate(times,1)]+[["Total",1278,678,*([78]*12)]],columns=["time","Total (PCU)","Total (คัน)",*vehicle_classes])
    composition=pd.DataFrame(
        [{"vehicle_class":code,"สัดส่วน (%)":1/12} for code in vehicle_classes]
        +[{"vehicle_class":"Total","สัดส่วน (%)":1.0}]
    )
    peak=pd.DataFrame([{"peak_mode":"rolling_60min","am_peak_search_window":"06:00-10:00","pm_peak_search_window":"15:00-19:00","am_peak_start":"07:00","am_peak_end":"08:00","am_peak_pcu":999,"am_phf":0.91,"pm_peak_start":"17:00","pm_peak_end":"18:00","pm_peak_pcu":1111,"pm_phf":0.92,"peak_selection_source":"user_confirmed"}])
    sheets={"Export_Metadata":pd.DataFrame([("peak_selection_source","user_confirmed"),("am_peak_start","07:00"),("am_peak_end","08:00"),("pm_peak_start","17:00"),("pm_peak_end","18:00")],columns=["field","value"]),"Peak_PHF":peak,"Hourly_Movement_PCU":hourly,"Hourly_Vehicle_Class":vehicle,"Mapping":mapping_frame,"Normalized_Data":normalized_frame,"Movement_Summary":movement_summary,"Setup":pd.DataFrame([("peak_selection_source","user_confirmed")],columns=["field","value"])}
    data={"sheets":sheets,"hourly_movement_pcu":hourly,"diagram_movement_codes":("NU","NW","NS","NE","SW","SN","SE","SU","WN","WE","WS","WU","EU","EN","EW","ES")}
    chart=_native_chart_source_data(hourly,composition)
    data["vehicle_composition_report"]=composition
    metadata={"report_title":"OOXML fidelity report","project_name":"Qualification","survey_point":"TMC-01","survey_date_text":"2026-09-27","weather":"Clear","responsible_party":"Test","survey_period":"07:00-19:00","am_peak_start":"07:00","am_peak_end":"08:00","pm_peak_start":"17:00","pm_peak_end":"18:00","peak_selection_source":"user_confirmed","north_label":"North","south_label":"South","east_label":"East","west_label":"West","north_road":"Road N","south_road":"Road S","east_road":"Road E","west_road":"Road W","caption_text":"Counts per 12 hours"}
    return data,metadata,chart


def _effective_payload(am_start: str, pm_start: str):
    data, metadata, chart = _payload()
    hourly = data["hourly_movement_pcu"]
    for period, start in (("am", am_start), ("pm", pm_start)):
        end = f"{int(start[:2])+1:02}:00"
        metadata[f"{period}_peak_start"] = start
        metadata[f"{period}_peak_end"] = end
        peak = data["sheets"]["Peak_PHF"]
        peak.loc[0, f"{period}_peak_start"] = start
        peak.loc[0, f"{period}_peak_end"] = end
        selected = hourly.loc[hourly["time"] == f"{start}-{end}"].iloc[0]
        peak.loc[0, f"{period}_peak_pcu"] = sum(selected[code] for code in MOVEMENT_CODES)
        export = data["sheets"]["Export_Metadata"]
        for key, value in ((f"{period}_peak_start", start), (f"{period}_peak_end", end)):
            export.loc[export["field"] == key, "value"] = value
    return data, metadata, chart


@pytest.mark.parametrize("am,pm,am_row,pm_row", [
    ("08:00", "17:00", 11, 20),
    ("07:00", "16:00", 10, 19),
    ("08:00", "16:00", 11, 19),
    ("17:00", "17:00", 20, 20),
    ("07:00", "17:00", 10, 20),
])
def test_effective_peak_binds_all_native_formulas_and_caches(tmp_path: Path, am: str, pm: str, am_row: int, pm_row: int) -> None:
    data, metadata, chart = _effective_payload(am, pm)
    mapping = load_template_map(DEFAULT_TEMPLATE_MAP_PATH)
    plan = resolve_template_write_plan(DEFAULT_TEMPLATE_PATH, mapping, data, metadata, chart)
    assert [(peak.worksheet_row, peak.helper_cell, peak.helper_value) for peak in plan.peak_formula_rows] == [
        (am_row, f"U{am_row}", am_row-8), (pm_row, f"U{pm_row}", pm_row-8)
    ]
    assert len(plan.formula_bindings) == 32
    assert not any(write.cell.startswith("U") for write in plan.summary_writes)
    output = tmp_path / "bound.xlsx"
    export_template_ooxml(DEFAULT_TEMPLATE_PATH, output, mapping, data, metadata, chart)
    verify_ooxml_against_plan(DEFAULT_TEMPLATE_PATH, output, mapping, plan)
    formulas = load_workbook(output, read_only=True, data_only=False)
    values = load_workbook(output, read_only=True, data_only=True)
    try:
        summary, cached = formulas["Summary"], values["Summary"]
        assert [summary[f"U{row}"].value for row in range(9, 23)] == list(range(1, 15))
        assert summary["A3"].value == f"Peak used for export: AM {am}-{metadata['am_peak_end']}; PM {pm}-{metadata['pm_peak_end']} (user_confirmed)"
        assert values["Peak_PHF"]["D2"].value == am
        assert values["Peak_PHF"]["H2"].value == pm
        fields = {row[0]: row[1] for row in values["Export_Metadata"].iter_rows(min_row=2, values_only=True)}
        assert fields["am_peak_start"] == am and fields["pm_peak_start"] == pm
        for code, slot in mapping["movement_diagram_cells"]["diagram_movements"].items():
            column = mapping["hourly_movement_table"]["columns"][code]
            for period, row in (("am", am_row), ("pm", pm_row)):
                ref = slot[f"{period}_peak_hour_cell"]
                assert f"$U${row},FALSE)" in summary[ref].value
                assert cached[ref].value == cached[f"{column}{row}"].value
            assert "$U$22,FALSE)" in summary[slot["total_12_hour_cell"]].value
            assert cached[slot["total_12_hour_cell"]].value == cached[f"{column}22"].value
            diagram_row = list(data["diagram_movement_codes"]).index(code)+2
            assert values["Diagram_Data"].cell(diagram_row, 3).value == cached[slot["pm_peak_hour_cell"]].value
            assert values["Diagram_Data"].cell(diagram_row, 4).value == cached[slot["am_peak_hour_cell"]].value
    finally:
        formulas.close(); values.close()


def test_confirmed_peak_change_rebinds_formulas_without_changing_helpers(tmp_path: Path) -> None:
    mapping = load_template_map(DEFAULT_TEMPLATE_MAP_PATH)
    outputs = []
    for start in ("07:00", "08:00"):
        data, metadata, chart = _effective_payload(start, "17:00")
        output = tmp_path / f"confirmed_{start[:2]}.xlsx"
        export_template_ooxml(DEFAULT_TEMPLATE_PATH, output, mapping, data, metadata, chart)
        formulas = load_workbook(output, read_only=True, data_only=False)
        values = load_workbook(output, read_only=True, data_only=True)
        try:
            outputs.append((formulas["Summary"]["J16"].value, values["Summary"]["J16"].value,
                [values["Summary"][f"U{row}"].value for row in range(9,23)],
                values["Peak_PHF"]["D2"].value, values["Diagram_Data"]["D2"].value,
                values["Summary"]["A3"].value))
        finally:
            formulas.close(); values.close()
    first, second = outputs
    assert "$U$10,FALSE)" in first[0] and "$U$11,FALSE)" in second[0]
    assert first[1] != second[1]
    assert first[2] == second[2] == list(range(1,15))
    assert first[3] == "07:00" and second[3] == "08:00"
    assert first[4] != second[4]
    assert "07:00-08:00" in first[5] and "08:00-09:00" in second[5]


def test_unmatched_effective_peak_fails_before_transport() -> None:
    data, metadata, chart = _payload()
    metadata["am_peak_start"], metadata["am_peak_end"] = "06:00", "07:00"
    with pytest.raises(ValueError, match="must match exactly one Summary time row"):
        resolve_template_write_plan(DEFAULT_TEMPLATE_PATH, load_template_map(DEFAULT_TEMPLATE_MAP_PATH), data, metadata, chart)


def test_unknown_native_hlookup_formula_fails_before_transport(tmp_path: Path) -> None:
    changed = tmp_path / "changed_template.xlsx"
    with zipfile.ZipFile(DEFAULT_TEMPLATE_PATH) as source, zipfile.ZipFile(changed, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(payload)
                formula = root.find(f".//{{{MAIN}}}c[@r='J16']/{{{MAIN}}}f")
                assert formula is not None
                formula.text = "HLOOKUP(J13,$W$9:$AL$22,$U$11,TRUE)"
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, payload)
    data, metadata, chart = _payload()
    with pytest.raises(ValueError, match="Unknown native HLOOKUP formula"):
        resolve_template_write_plan(changed, load_template_map(DEFAULT_TEMPLATE_MAP_PATH), data, metadata, chart)


def test_shared_com_contract_matches_ooxml_package(tmp_path: Path) -> None:
    data, metadata, chart = _payload()
    mapping = load_template_map(DEFAULT_TEMPLATE_MAP_PATH)
    plan = resolve_template_write_plan(DEFAULT_TEMPLATE_PATH, mapping, data, metadata, chart)
    output = tmp_path / "shared_write_contract.xlsx"
    export_template_ooxml(DEFAULT_TEMPLATE_PATH, output, mapping, data, metadata, chart)
    result = verify_ooxml_against_plan(DEFAULT_TEMPLATE_PATH, output, mapping, plan)
    assert "planned_summary_writes" in result
    assert len(plan.diagram_rows) == 16
    assert {item.code for item in plan.diagram_rows} == set(MOVEMENT_CODES)
    assert {name for name, _ in plan.support_sheets} >= {"Export_Metadata", "Peak_PHF", "Hourly_Movement_PCU", "Hourly_Vehicle_Class"}
    assert [write.order for write in plan.ordered_writes] == list(range(1, len(plan.ordered_writes) + 1))
    assert any(write.sheet == "Peak_PHF" and write.cell == "D2" and write.value_type == "text" for write in plan.ordered_writes)
    assert any(write.sheet == "Diagram_Data" and write.cell == "B2" and write.value_type == "formula" for write in plan.ordered_writes)


def test_com_transport_receives_exact_shared_summary_intents(monkeypatch) -> None:
    data, metadata, chart = _payload()
    mapping = load_template_map(DEFAULT_TEMPLATE_MAP_PATH)
    plan = resolve_template_write_plan(DEFAULT_TEMPLATE_PATH, mapping, data, metadata, chart)
    received = []
    def capture(worksheet, cell, value, guard, source, force=False):
        received.append((cell, value, source, force))
    monkeypatch.setattr(excel_com_export, "_set_cell", capture)
    excel_com_export._apply_summary_write_plan(object(), plan, object())
    assert received == [(write.cell, write.value, write.source, write.force) for write in plan.summary_writes]
    cells = {binding.cell: SimpleNamespace(Formula=binding.old_formula) for binding in plan.formula_bindings}
    worksheet = SimpleNamespace(Range=lambda ref: cells[ref])
    diagnostics = excel_com_export.ExcelComExportDiagnostics()
    excel_com_export._apply_peak_formula_bindings(worksheet, plan, diagnostics)
    assert [cells[binding.cell].Formula for binding in plan.formula_bindings] == [binding.new_formula for binding in plan.formula_bindings]
    assert len(diagnostics.formula_cells_overwritten) == 32


@pytest.mark.parametrize("damage", ["skip_write", "wrong_destination", "swap_columns", "missing_support", "stale_chart_cache", "stale_movement_cache", "stale_vehicle_total_cache", "missing_hourly_total_cache", "wrong_hourly_total_cache", "missing_grand_total_cache", "missing_rollup_cache", "missing_title_cache", "stale_text_date_cache", "wrong_peak_binding", "changed_u_helper", "changed_total_formula", "changed_formula", "changed_chart_definition", "changed_drawing"])
def test_write_contract_verifier_rejects_negative_cases(tmp_path: Path, damage: str) -> None:
    data, metadata, chart = _payload()
    mapping = load_template_map(DEFAULT_TEMPLATE_MAP_PATH)
    plan = resolve_template_write_plan(DEFAULT_TEMPLATE_PATH, mapping, data, metadata, chart)
    clean = tmp_path / "clean.xlsx"
    broken = tmp_path / f"broken_{damage}.xlsx"
    export_template_ooxml(DEFAULT_TEMPLATE_PATH, clean, mapping, data, metadata, chart)
    with zipfile.ZipFile(clean) as source, zipfile.ZipFile(broken, "w") as target:
        summary_part = "xl/worksheets/sheet1.xml"
        chart_part = mapping["chart_anchors"]["vehicle_composition_bar_chart"]["chart_xml"]
        for info in source.infolist():
            part = info.filename
            if damage == "missing_support" and part == "xl/worksheets/sheet2.xml":
                continue
            payload = source.read(part)
            if part == summary_part and damage in {"skip_write", "wrong_destination", "swap_columns", "stale_movement_cache", "stale_vehicle_total_cache", "missing_hourly_total_cache", "wrong_hourly_total_cache", "missing_grand_total_cache", "missing_rollup_cache", "missing_title_cache", "stale_text_date_cache", "wrong_peak_binding", "changed_u_helper", "changed_total_formula", "changed_formula"}:
                root = ET.fromstring(payload)
                cells = {cell.attrib["r"]: cell for cell in root.findall(f".//{{{MAIN}}}c")}
                if damage == "skip_write":
                    row = root.find(f".//{{{MAIN}}}row[@r='10']")
                    row.remove(cells["W10"])
                elif damage == "wrong_destination":
                    cells["W10"].find(f"{{{MAIN}}}v").text = "999"
                elif damage == "swap_columns":
                    first = cells["W10"].find(f"{{{MAIN}}}v")
                    second = cells["X10"].find(f"{{{MAIN}}}v")
                    first.text, second.text = second.text, first.text
                elif damage == "stale_movement_cache":
                    cells["M16"].find(f"{{{MAIN}}}v").text = "999"
                elif damage == "stale_vehicle_total_cache":
                    cells["X40"].find(f"{{{MAIN}}}v").text = "999"
                elif damage == "missing_hourly_total_cache":
                    cells["W22"].find(f"{{{MAIN}}}v").text = None
                elif damage == "wrong_hourly_total_cache":
                    cells["W22"].find(f"{{{MAIN}}}v").text = "999"
                elif damage == "missing_grand_total_cache":
                    cells["AM22"].find(f"{{{MAIN}}}v").text = None
                elif damage == "missing_rollup_cache":
                    cells["F30"].find(f"{{{MAIN}}}v").text = None
                elif damage == "missing_title_cache":
                    cells["C8"].find(f"{{{MAIN}}}v").text = None
                elif damage == "stale_text_date_cache":
                    cells["C9"].find(f"{{{MAIN}}}v").text = "999"
                elif damage == "wrong_peak_binding":
                    cells["M16"].find(f"{{{MAIN}}}f").text = "HLOOKUP(M13,$W$9:$AL$22,$U$11,FALSE)"
                elif damage == "changed_u_helper":
                    cells["U11"].find(f"{{{MAIN}}}v").text = "2"
                elif damage == "changed_total_formula":
                    cells["W22"].find(f"{{{MAIN}}}f").text = "SUM(W10:W20)"
                else:
                    cells["M16"].find(f"{{{MAIN}}}f").text = "SUM(1)"
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            if part == chart_part and damage == "stale_chart_cache":
                root = ET.fromstring(payload)
                point = root.find(f".//{{{CHART}}}numCache/{{{CHART}}}pt/{{{CHART}}}v")
                assert point is not None
                point.text = "999"
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            if part == chart_part and damage == "changed_chart_definition":
                root = ET.fromstring(payload)
                formula = root.find(f".//{{{CHART}}}f")
                assert formula is not None
                formula.text = "Summary!$A$1:$A$12"
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            if part == "xl/drawings/drawing1.xml" and damage == "changed_drawing":
                payload += b"<!-- unintended drawing mutation -->"
            target.writestr(info, payload)
    with pytest.raises(AssertionError, match="write-plan verification failed"):
        verify_ooxml_against_plan(DEFAULT_TEMPLATE_PATH, broken, mapping, plan)


def _parts(path):
    with zipfile.ZipFile(path) as z: return {n:z.read(n) for n in z.namelist()}

def _shape_count(drawing_xml):
    root=ET.fromstring(drawing_xml)
    return len(root.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}sp"))

def _chart_summary(xml):
    root=ET.fromstring(xml); out=[]
    for ser in root.findall(".//c:ser",{"c":CHART}):
        formulas=[n.text for n in ser.findall(".//c:f",{"c":CHART})]
        caches=[]
        for kind in ("cat","val"):
            box=ser.find(f"c:{kind}",{"c":CHART}); cache=box.find(".//c:strCache" if kind=="cat" else ".//c:numCache",{"c":CHART})
            caches.append([pt.findtext("c:v",namespaces={"c":CHART}) for pt in cache.findall("c:pt",{"c":CHART})] if cache is not None else [])
        out.append((formulas,caches))
    return out


def test_ooxml_export_preserves_template_and_adds_full_report_surfaces(tmp_path: Path) -> None:
    source=DEFAULT_TEMPLATE_PATH; output=tmp_path/"native_report.xlsx"; before_sha=hashlib.sha256(source.read_bytes()).hexdigest(); before=_parts(source)
    data,metadata,chart=_payload(); mapping=load_template_map(DEFAULT_TEMPLATE_MAP_PATH)
    plan=resolve_template_write_plan(source,mapping,data,metadata,chart)
    result=export_template_ooxml(source,output,mapping,data,metadata,chart)
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before_sha
    after=_parts(output)
    assert zipfile.is_zipfile(output)
    with zipfile.ZipFile(output) as z:
        assert z.testzip() is None
        for name,payload in after.items():
            if name.endswith(".xml") or name.endswith(".rels") or name=="[Content_Types].xml": ET.fromstring(payload)
        assert not any(name.startswith("xl/externalLinks/") for name in after)
        assert b"#REF!" not in b"".join(payload for name,payload in after.items() if name.endswith(".xml"))
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore",message="DrawingML support is incomplete.*",category=UserWarning)
        formula_book=load_workbook(output,data_only=False)
        value_book=load_workbook(output,data_only=True,read_only=True)
    try:
        assert formula_book.sheetnames==["Summary",*data["sheets"],"Diagram_Data"]
        assert formula_book["Summary"]["B2"].value==metadata["report_title"]
        assert formula_book["Summary"]["E5"].value==metadata["survey_point"]
        assert formula_book["Summary"]["A3"].value=="Peak used for export: AM 07:00-08:00; PM 17:00-18:00 (user_confirmed)"
        assert formula_book["Summary"]["E4"].value==metadata["project_name"]
        assert formula_book["Summary"]["V10"].value=="07:00-08:00"
        assert formula_book["Export_Metadata"]["A2"].value=="peak_selection_source"
        assert formula_book["Export_Metadata"]["B2"].value=="user_confirmed"
        assert formula_book["Peak_PHF"]["E2"].value=="08:00"
        assert formula_book["Peak_PHF"]["L2"].value=="user_confirmed"
        assert formula_book["Diagram_Data"]["A2"].value=="NU"
        assert formula_book["Diagram_Data"]["B2"].value.startswith("=IFERROR(INDEX(")
        assert [formula_book["Summary"][f"U{row}"].value for row in range(9,23)]==list(range(1,15))
        assert formula_book["Summary"]["U11"].value==3
        assert formula_book["Summary"]["U20"].value==12
        assert formula_book["Summary"]["U22"].value==14
        assert formula_book["Summary"]["M16"].value=="=HLOOKUP(M13,$W$9:$AL$22,$U$10,FALSE)"
        assert value_book["Summary"]["M16"].value==1
        assert value_book["Summary"]["M15"].value==11
        assert value_book["Summary"]["M14"].value==78
        assert value_book["Summary"]["F30"].value==136
        assert value_book["Summary"]["F31"].value==296
        assert value_book["Summary"]["F32"].value==2688
        assert value_book["Summary"]["C8"].value==metadata["survey_point"]
        assert value_book["Summary"]["W22"].value==78
        assert value_book["Summary"]["AM22"].value==2688
        for offset in range(16):
            ref=formula_book["Summary"].cell(22,23+offset).coordinate
            assert formula_book["Summary"][ref].value==f"=SUM({ref[:-2]}10:{ref[:-2]}21)"
            assert value_book["Summary"][ref].value==78+12*offset
        assert formula_book["Export_Metadata"]["A1"].style_id>=160
        assert formula_book["Summary"]["E5"].style_id==load_workbook(source,data_only=False)["Summary"]["E5"].style_id
        # The vehicle table uses its populated support-sheet frame when the
        # compatibility payload key is omitted.
        assert formula_book["Summary"]["X28"].value==51
        assert formula_book["Summary"]["Y28"].value==1
        assert formula_book["Summary"]["AJ39"].value==12
        assert formula_book["Summary"]["X27"].value=="Total (คัน)"
        for row in range(28,40):
            for column in range(24,37):
                assert isinstance(formula_book["Summary"].cell(row,column).value,(int,float))
                support_value=formula_book["Hourly_Vehicle_Class"].cell(row-26,column-21).value
                assert formula_book["Summary"].cell(row,column).value==support_value
        assert formula_book["Summary"]["X40"].value=="=SUM(X28:X39)"
        assert value_book["Summary"]["X40"].value==678
        code_order=list(data["hourly_movement_pcu"].columns[1:-1])
        for code,info in mapping["movement_diagram_cells"]["diagram_movements"].items():
            code_index=code_order.index(code)
            total=sum(i+code_index for i in range(1,13))
            movement_index=data["diagram_movement_codes"].index(code)+2
            assert value_book["Diagram_Data"].cell(movement_index,2).value==total
            assert value_book["Diagram_Data"].cell(movement_index,3).value==11+code_index
            assert value_book["Diagram_Data"].cell(movement_index,4).value==1+code_index
            assert value_book["Summary"][info["total_12_hour_cell"]].value==total
            assert value_book["Summary"][info["pm_peak_hour_cell"]].value==11+code_index
            assert value_book["Summary"][info["am_peak_hour_cell"]].value==1+code_index
    finally:
        formula_book.close(); value_book.close()

    original_formula_parts,original_formulas=iter_formula_cells(source); output_sheets,output_formulas=iter_formula_cells(output)
    old_summary={cell:formula for sheet,cell,formula in original_formulas if sheet=="Summary" and cell!="B2"}
    new_summary={cell:formula for sheet,cell,formula in output_formulas if sheet=="Summary"}
    bound={binding.cell:binding.new_formula[1:] for binding in plan.formula_bindings if binding.cell in old_summary}
    assert new_summary=={**old_summary,**bound}
    assert set(data["sheets"])|{"Summary","Diagram_Data"}==set(output_sheets)
    assert not audit_template(output).warning_lines()

    # Protected formula nodes are exactly preserved, including formula text.
    protected=set()
    from openpyxl.utils.cell import range_boundaries
    for area in mapping["protected_formula_ranges"]:
        c1,r1,c2,r2=range_boundaries(area)
        for row in range(r1,r2+1):
            for col in range(c1,c2+1):
                from openpyxl.utils.cell import get_column_letter
                protected.add(f"{get_column_letter(col)}{row}")
    assert {c:f for c,f in old_summary.items() if c in protected and c not in bound}=={c:f for c,f in new_summary.items() if c in protected and c not in bound}
    assert result.formula_cells_overwritten==("Summary!B2",)
    assert "Summary!B2" not in result.formula_cells_preserved

    # Original Summary part, formulas, merges, and template styling remain intact.
    original_wb=load_workbook(source,data_only=False)
    output_wb=load_workbook(output,data_only=False)
    try:
        assert {str(r) for r in original_wb["Summary"].merged_cells.ranges}=={str(r) for r in output_wb["Summary"].merged_cells.ranges}
        for ref in ("E4","E5","K5","Q5","E6","K6","G12","K11","Q17","R19","D24","D26","O32","K32","E35"):
            assert original_wb["Summary"][ref].style_id==output_wb["Summary"][ref].style_id
    finally: original_wb.close(); output_wb.close()

    # Existing relationship targets and IDs remain stable; new IDs are unique.
    with zipfile.ZipFile(source) as oldzip,zipfile.ZipFile(output) as newzip:
        oldrels=ET.fromstring(oldzip.read("xl/_rels/workbook.xml.rels")); rels=ET.fromstring(newzip.read("xl/_rels/workbook.xml.rels"))
        old_by_id={r.attrib["Id"]:r.attrib["Target"] for r in oldrels}; all_ids=[r.attrib["Id"] for r in rels]
        assert len(all_ids)==len(set(all_ids))
        new_by_id={r.attrib["Id"]:r.attrib["Target"] for r in rels}
        for rid,target in old_by_id.items():
            if rid!="rId5": assert new_by_id[rid]==target
        wb=ET.fromstring(newzip.read("xl/workbook.xml")); sheet_nodes=wb.findall(".//m:sheet",{"m":MAIN})
        assert all(node.get("state","visible")=="visible" for node in sheet_nodes)
        assert sheet_nodes[0].attrib["name"]=="Summary" and sheet_nodes[0].attrib["sheetId"]=="1" and sheet_nodes[0].attrib[f"{{{DOC_REL}}}id"]=="rId1"
        types=ET.fromstring(newzip.read("[Content_Types].xml")); overrides={o.attrib["PartName"] for o in types.findall("{*}Override")}
        assert all("/"+part in overrides for part in result.added_parts)
        assert "xl/calcChain.xml" not in newzip.namelist()
        calc=wb.find("m:calcPr",{"m":MAIN}); assert calc is not None
        assert calc.attrib=={"calcId":calc.attrib.get("calcId"),"calcMode":"auto","fullCalcOnLoad":"1","forceFullCalc":"1"}
        chart1=_chart_summary(newzip.read("xl/charts/chart1.xml")); chart2=_chart_summary(newzip.read("xl/charts/chart2.xml"))
        assert chart1[0][0]==["Summary!$V$10:$V$21","Summary!$AM$10:$AM$21"]
        assert chart2[0][0]==["Summary!$AP$39:$BA$39","Summary!$AP$40:$BA$40"]
        assert chart1[0][1]==[chart["hourly_pcu"]["categories"],[str(v) for v in chart["hourly_pcu"]["values"]]]
        assert chart2[0][1]==[chart["vehicle_composition"]["categories"],[str(v) for v in chart["vehicle_composition"]["values"]]]
        assert chart["vehicle_composition"]["values"]==[1/12]*12
        assert len(chart2[0][1][0])==len(chart2[0][1][1])==12
        assert any(float(value)>0 for value in chart2[0][1][1])
        summary_xml=ET.fromstring(newzip.read("xl/worksheets/sheet1.xml"))
        summary_cells={cell.attrib["r"]:cell for cell in summary_xml.findall(".//m:sheetData/m:row/m:c",{"m":MAIN})}
        for address in ("W28","W40","X40","Y40","AP38","AP40","BA40"):
            cell=summary_cells[address]
            assert cell.find("m:f",{"m":MAIN}) is not None
            assert cell.get("t")!="e"
            value=cell.find("m:v",{"m":MAIN})
            assert value is not None and value.text
        assert oldzip.read("xl/drawings/drawing1.xml")==newzip.read("xl/drawings/drawing1.xml")
        assert _shape_count(newzip.read("xl/drawings/drawing1.xml"))==33
        for name in ("xl/charts/style1.xml","xl/charts/colors1.xml","xl/theme/theme1.xml","xl/worksheets/_rels/sheet1.xml.rels","xl/drawings/_rels/drawing1.xml.rels"):
            assert oldzip.read(name)==newzip.read(name)

    manifest=compare_package_parts(source,output,expected_changed=result.changed_parts,expected_added=result.added_parts,expected_removed=result.removed_parts)
    assert all(not entry.classification.startswith("UNEXPECTED_") for entry in manifest)
    assert {entry.classification for entry in manifest}=={"UNCHANGED_EXPECTED","CHANGED_EXPECTED","ADDED_EXPECTED","REMOVED_EXPECTED"}
    diff_summary=summarize_package_diff(manifest)
    assert diff_summary["charts"]["changed expected"]==2
    assert diff_summary["drawings"]["unchanged"]>=1
    # XML part inventory stays deterministic.
    assert inventory_package(output)==inventory_package(output)


def test_ooxml_export_rejects_unmapped_and_protected_formula_writes() -> None:
    mapping=load_template_map(DEFAULT_TEMPLATE_MAP_PATH)
    doc=ET.Element("root")
    from xml.dom import minidom
    summary=minidom.parseString(f'<worksheet xmlns="{MAIN}"><sheetData><row r="2"><c r="E5"><f>1+1</f><v>2</v></c><c r="D200"><v>0</v></c></row></sheetData></worksheet>')
    cells={c.getAttribute("r"):c for c in summary.getElementsByTagNameNS(MAIN,"c")}
    import pytest
    with pytest.raises(ValueError,match="Unsafe OOXML write"):
        _put_summary(summary,cells,"D200","unsafe",mapping,intent="test")
    with pytest.raises(ValueError,match="formula overwrite"):
        _put_summary(summary,cells,"E5","unsafe",mapping,intent="test")
